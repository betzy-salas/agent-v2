#!/usr/bin/env python3
"""Las dos láminas para correo (scripts/13_laminas_correo.py) en PowerPoint (.pptx), con el diseño del deck corporativo (16:9).

Lámina 1 · Resultado T5: ranking (puntaje, barra, nivel) y, por vendor, lo que ofrece según sus tres primeras fortalezas verificadas.
Lámina 2 · Implementación y despliegue: ficha de implementación (enfoque, duración, implementador, migración, hypercare, despliegue)
y lo que propone cada vendor para implantar (dos primeros pros del informe de implementación).

Mismas fuentes y mismos textos que el script 13; no calcula ni redacta nada.
Uso:  python scripts/14_laminas_correo_pptx.py [--base carpeta]      → <base>/LAMINAS_CORREO.pptx      Requiere: pip install python-pptx
"""
import datetime, importlib.util, sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent; RAIZ = AQUI.parent
def _cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, AQUI / archivo); mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod
p12 = _cargar('p12', '12_presentacion_t5_pptx.py'); p13 = _cargar('p13', '13_laminas_correo.py'); p10 = p12.p10; m = p12.m
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W, H = p12.W, p12.H; NAVY, BLUE, NAVY_T, ORANGE, INK, GREY, LBLUE = p12.NAVY, p12.BLUE, p12.NAVY_T, p12.ORANGE, p12.INK, p12.GREY, p12.LBLUE
rect, texto, tabla, barra, rgb = p12.rect, p12.texto, p12.tabla, p12.barra, p12.rgb
LEVEL_FILL, PILL = p12.LEVEL_FILL, p12.PILL

def pie(slide, fuente, logo):
    texto(slide, 0.42, 7.1, 7.5, 0.3, [[(fuente, {'color': GREY, 'size': 6.8})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 4.2, 7.18, 5, 0.3, [[('Datos elaborados por BCP para uso Interno', {'color': LBLUE, 'size': 6.8})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    if logo.exists(): slide.shapes.add_picture(str(logo), Inches(W - 0.42 - 0.62), Inches(7.12), Inches(0.62))

def h3(slide, x, y, w, txt, size=15):
    texto(slide, x, y, w, 0.4, [[(txt, {'bold': True, 'color': BLUE, 'size': size})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)

def construir(base, salida):
    V = p10.cargar(base)
    if not V: sys.exit(f'No encontré ningún 9_semaforo_T5.md bajo {base}.')
    for i, d in enumerate(V):
        d['_i'] = i; d['dist_cat'], d['dist'] = p10.distribucion_items(d['ruta']); d['ficha'] = p10.leer_ficha(d['ruta'].parent / '10_ficha_T5.md')
    orden = sorted(V, key=lambda d: -(d['pct'] or 0)); pos = {d['vendor']: k + 1 for k, d in enumerate(orden)}
    logo = base / 'assets' / 'logo.png'
    MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    t = datetime.date.today(); mes = f'{MESES[t.month - 1].capitalize()} {t.year}'
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H); blank = prs.slide_layouts[6]

    # ---- lámina 1 · resultado T5
    s = prs.slides.add_slide(blank); barra(s, 'Revisión Técnica Resultados T5 · Resultado por vendor', 'Capacidades de desarrollo (hoja T5) · ' + mes)
    h3(s, 0.42, 0.78, 5.0, 'Ranking por puntaje T5')
    y = 1.3
    for d in orden:
        c = p10.color_de(d['pct']); fill, ink = PILL[c]
        rect(s, 0.42, y + 0.14, 0.34, 0.34, ORANGE, shape=MSO_SHAPE.OVAL)
        texto(s, 0.42, y + 0.14, 0.34, 0.34, [[(str(pos[d['vendor']]), {'bold': True, 'color': 'FFFFFF', 'size': 10})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
        texto(s, 0.86, y, 1.55, 0.62, [[(m.nombre(d['vendor']), {'bold': True, 'color': NAVY_T, 'size': 11.5})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=True, min_size=9)
        texto(s, 2.42, y, 0.75, 0.62, [[(str(d['pct']), {'bold': True, 'color': ink, 'size': 16, 'font': p12.MONO})]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
        rect(s, 3.28, y + 0.24, 1.3, 0.14, 'EEF1F4', shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, 3.28, y + 0.24, max(0.14, 1.3 * (d['pct'] or 0) / 100), 0.14, LEVEL_FILL[c], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, 4.7, y + 0.17, 1.12, 0.28, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        texto(s, 4.7, y + 0.17, 1.12, 0.28, [[(d['nivel'], {'bold': True, 'color': ink, 'size': 7.5})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
        ln = s.shapes.add_connector(1, Inches(0.42), Inches(y + 0.66), Inches(5.82), Inches(y + 0.66)); ln.line.color.rgb = rgb('E3E6EC'); ln.line.width = Pt(0.75)
        y += 0.72
    xr, wr = 6.1, 6.8
    h3(s, xr, 0.78, wr, 'Lo que ofrece cada propuesta y sustenta su resultado')
    y = 1.2; alto = (6.3 - y) / len(orden)
    for d in orden:
        f = d['ficha'] or {}; paras = [[(f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])}', {'bold': True, 'color': NAVY_T, 'size': 8.6})]]
        for x in (f.get('fortalezas') or [])[:3]:
            tit, que, items = p13.lo_que_ofrece(x)
            if not que: continue
            runs = [(tit + ': ', {'bold': True, 'color': NAVY_T})] if tit else []
            runs.append((m.nombres_en_texto(que), {}))
            if items: runs.append((f' ({items})', {'color': GREY, 'size': 6.4}))
            paras.append(runs)
        texto(s, xr, y, wr, alto - 0.04, paras, size=7.4, min_size=6, espacio=1)
        ln = s.shapes.add_connector(1, Inches(xr), Inches(y + alto - 0.03), Inches(xr + wr), Inches(y + alto - 0.03)); ln.line.color.rgb = rgb('E3E6EC'); ln.line.width = Pt(0.75)
        y += alto
    rect(s, 0.42, 6.42, 12.5, 0.58, 'EEF1F4', shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    texto(s, 0.52, 6.42, 12.3, 0.58, [[('Cómo leer. ', {'bold': True, 'color': NAVY_T}), ('Puntaje T5 = media de las calificaciones del Banco (0–4) en la hoja T5 ponderada por criticidad (Indispensable y Killer ×3, Deseable ×1), sobre 100. Niveles: Cumple plenamente ≥ 85 · Cumple 70–84 · Parcial 50–69 · No cumple < 50. Las calificaciones son las del Banco y no se modifican; la verificación documental solo aporta la evidencia que sustenta fortalezas, debilidades y contradicciones.', {})]], size=7.8, color='3D4A5C', anchor=MSO_ANCHOR.MIDDLE, min_size=6.5)
    pie(s, 'Fuente: Matriz técnica del Banco, hoja T5 · semáforo y ficha T5 de cada vendor', logo)

    # ---- lámina 2 · implementación y despliegue
    IMP = p10.datos_implementacion(V); g = IMP['general']
    if g:
        DESP, PLAT = p10.datos_despliegue(orden); desp = {v: filas for v, _, filas in DESP}
        s = prs.slides.add_slide(blank); barra(s, 'Revisión Técnica Resultados T5 · Implementación y despliegue', 'Cómo propone implantar cada vendor y modelos de despliegue · ' + mes)
        filas = []
        for d in orden:
            x = IMP['vendors'].get(d['vendor'])
            if not x: continue
            fi = {e: q for e, q, _, _ in p10.ficha_impl_filas(x)}; fl = desp.get(d['vendor'], [])
            ofrec = [nm for nm, cls, _, _ in fl if cls == 'si']; cond = [nm for nm, cls, _, _ in fl if cls == 'cond']
            nubes = []
            for _, _, _, marcas, *_ in fl:
                for pn, est in marcas.items():
                    if est == 'of' and pn not in nubes and pn != 'Nube privada': nubes.append(pn)
            dcel = [('Ofrece: ', {'bold': True}), (', '.join(ofrec) if ofrec else '—', {})]
            if cond: dcel += [(' · Condicionado: ', {'bold': True}), (', '.join(cond), {})]
            dcel += [('\nNubes: ', {'bold': True}), (', '.join(nubes) if nubes else 'no especifica', {})]
            celda = lambda k: m.nombres_en_texto(fi.get(k, '—'))
            filas.append([f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])}', celda('Enfoque'), celda('Duración total'), celda('Implementador'), celda('Migración de datos'), celda('Hypercare'), dcel])
        col_w = [1.25, 2.1, 1.55, 1.9, 2.15, 1.55, 2.0]; sz = 7.0
        shp = tabla(s, 0.42, 0.8, 12.5, col_w, ['Vendor (ranking T5)', 'Enfoque de implementación', 'Duración propuesta', 'Implementador', 'Migración de datos', 'Hypercare', 'Modelos de despliegue y nubes ofrecidos'], filas, size=sz, head_size=7.4)
        # alto de cada fila según el texto más largo de la fila (PowerPoint expande las filas; así lo siguiente se coloca donde termina la tabla)
        t = shp.table; t.rows[0].height = Inches(0.34); y_fin = 0.8 + 0.34
        for i, fila in enumerate(filas, 1):
            lineas = max(p12.estimar_lineas([[(c, {})] if isinstance(c, str) else c], col_w[j] - 0.12, sz) + (0 if isinstance(c, str) else 1) for j, c in enumerate(fila))   # +1: la celda de despliegue tiene salto de línea
            h = lineas * sz * 1.2 / 72 + 0.08; t.rows[i].height = Inches(h); y_fin += h
        y0 = y_fin + 0.06
        texto(s, 0.42, y0, 12.5, 0.28, [[('Ficha de implementación de cada vendor, tomada de su propuesta y anexos. Despliegue: modelos que el vendor ofrece (Sí) o condiciona y nubes en firme para esos modelos, según sus documentos.', {'color': GREY, 'size': 7})]], ajustar=False)
        h3(s, 0.42, y0 + 0.26, 12.5, 'Lo que propone cada vendor para implantar', size=12)
        y_pros = y0 + 0.64; paso = min(0.4, (7.02 - y_pros) / 3)
        k = 0
        for d in orden:
            x = IMP['vendors'].get(d['vendor'])
            if not x: continue
            its = [m.nombres_en_texto(i['afirmacion']).rstrip('.') for i in (x.get('pros_items') or [])[:2]]
            if not its: continue
            col, fila = k % 2, k // 2; xx = 0.42 + col * 6.35; yy = y_pros + fila * paso
            texto(s, xx, yy, 6.15, paso - 0.02, [[(f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])}. ', {'bold': True, 'color': NAVY_T}), (' · '.join(its), {})]], size=7.4, min_size=6, espacio=0)
            k += 1
        pie(s, 'Fuente: informe de implementación (7_) y análisis de despliegue (8_) de cada vendor, sobre su propuesta y anexos', logo)
    prs.save(salida); return len(prs.slides)

def main():
    args = sys.argv[1:]
    base = Path(args[args.index('--base') + 1]) if '--base' in args else RAIZ / 'vendors'
    if not base.is_absolute(): base = RAIZ / base
    salida = base / 'LAMINAS_CORREO.pptx'; print(f'PPTX: {salida} · {construir(base, salida)} diapositivas')

if __name__ == '__main__':
    main()
