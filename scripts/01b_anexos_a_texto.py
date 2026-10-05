#!/usr/bin/env python3
"""Convierte los anexos no-PDF de docs/raw (xlsx, pptx) a docs/txt/<nombre>.txt
y genera docs/txt/<nombre>.toc.md.

Marcadores (compatibles con los del PDF):
  Excel : ===== SHEET <hoja> =====   y cada fila   f.<n>: celda | celda | ...
  PPTX  : ===== SLIDE <n> =====
Cita en los verificadores: «Anexo 4_Sunline, hoja T3, f.27» o «Key Roles, slide 2».

Se excluye la matriz de evaluación del Banco (el xlsx con la hoja de método o con hojas T1–T8 y
columna VA): ya está volcada en matriz/T*.txt. Usa --incluir-matriz para forzarla.

Uso: python scripts/01b_anexos_a_texto.py <vendor> [--incluir-matriz]
"""
import re, sys, zipfile, pathlib, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

from _comun import arg_vendor, vendor_dir, escribir_nombres, es_matriz_evaluacion, archivos_raw, slug_de, requiere
requiere('openpyxl')
v = arg_vendor(); d = vendor_dir(v)
raw, out = d / 'docs' / 'raw', d / 'docs' / 'txt'
out.mkdir(parents=True, exist_ok=True)
INCLUIR_MATRIZ = '--incluir-matriz' in sys.argv

def celda(c):
    if c is None: return ''
    s = str(c).replace('\r', ' ').replace('\n', ' ⏎ ').strip()
    return re.sub(r'\s+', ' ', s)

def xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    bloques, toc = [], []
    for ws in wb.worksheets:
        filas, n_datos, pista = [], 0, ''
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [celda(c) for c in row]
            while vals and vals[-1] == '': vals.pop()
            if not any(vals): continue
            n_datos += 1
            if not pista: pista = ' | '.join(v for v in vals if v)[:100]
            filas.append(f'f.{i}: ' + ' | '.join(vals))
        bloques.append(f'===== SHEET {ws.title} =====\n' + '\n'.join(filas))
        toc.append((ws.title, n_datos, pista))
    return bloques, toc

def pptx(path):
    z = zipfile.ZipFile(path)
    slides = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                    key=lambda s: int(re.findall(r'\d+', s)[0]))
    bloques, toc = [], []
    for k, s in enumerate(slides, 1):
        xml = z.read(s).decode('utf8')
        # un párrafo por <a:p>, texto de todos los <a:t>
        parrafos = [''.join(re.findall(r'<a:t>([^<]*)</a:t>', p)) for p in re.findall(r'<a:p>.*?</a:p>', xml, re.S)]
        parrafos = [re.sub(r'\s+', ' ', p).strip() for p in parrafos if p.strip()]
        notas = s.replace('slides/slide', 'notesSlides/notesSlide')
        if notas in z.namelist():
            nt = ' '.join(re.findall(r'<a:t>([^<]*)</a:t>', z.read(notas).decode('utf8'))).strip()
            if nt: parrafos.append('NOTAS: ' + nt)
        bloques.append(f'===== SLIDE {k} =====\n' + '\n'.join(parrafos))
        toc.append((f'slide {k}', len(parrafos), parrafos[0][:100] if parrafos else ''))
    return bloques, toc

def escribir(path, bloques, toc, unidad):
    stem = slug_de(path, raw)
    (out / f'{stem}.txt').write_text('\n\n'.join(bloques) + '\n', encoding='utf-8')
    lineas = [f'# Mapa de secciones — {stem}', '',
              f'Archivo texto: `docs/txt/{stem}.txt`. Origen: `docs/raw/{path.name}`.',
              f'Marcadores `===== {unidad} <nombre> =====`; en Excel cada fila va como `f.<n>: …` (n = número de fila en la hoja).',
              f'Cita: «{stem}, {unidad.lower()} <nombre>' + (', f.<n>»' if unidad == 'SHEET' else '»'), '',
              f'| {unidad.lower()} | filas/párrafos con datos | primer contenido |', '|---|---|---|']
    lineas += [f'| {t} | {n} | {p} |' for t, n, p in toc]
    (out / f'{stem}.toc.md').write_text('\n'.join(lineas) + '\n', encoding='utf-8')
    print(f'{path.relative_to(raw)}: {len(bloques)} {unidad.lower()}(s) → {stem}.txt / .toc.md')

for p in archivos_raw(raw):
    if p.suffix.lower() in ('.xlsx', '.xlsm'):
        if not INCLUIR_MATRIZ and es_matriz_evaluacion(p):
            print(f'{p.relative_to(raw)}: es la matriz de evaluación del Banco, omitida (ya está en matriz/)'); continue
        escribir(p, *xlsx(p), 'SHEET')
    elif p.suffix.lower() == '.pptx':
        escribir(p, *pptx(p), 'SLIDE')
    elif p.suffix.lower() != '.pdf':
        print(f'{p.relative_to(raw)}: formato no soportado, omitido', file=sys.stderr)

print('NOMBRES.md actualizado:', escribir_nombres(d).relative_to(d.parent.parent))
