#!/usr/bin/env python3
"""F1 · Rúbrica corporativa de implementación (xlsx del Banco con hojas «Criterios Estrat Implementación»,
«Criterios Fabricante», «Criterios Implementador») → vendors/<vendor>/matriz/IMPL.txt

Por cada criterio escribe: código, bloque, categoría, criterio, consideraciones, peso relativo, puntaje máximo,
escala (texto completo, 0–5), calificación del Banco si ya existe, y la respuesta del vendor tomada de su Anexo de
perfiles (hojas «T1 Fabricante | Perfil» y «T2 Implementadores | Perfil») ya convertido en docs/txt.

Uso: python scripts/02b_rubrica_implementacion.py <vendor> [ruta_al_xlsx]
"""
import re, sys, warnings
warnings.filterwarnings("ignore")
from pathlib import Path
from _comun import arg_vendor, vendor_dir, archivos_raw, es_rubrica_implementacion, requiere
requiere('openpyxl')
from openpyxl import load_workbook

v = arg_vendor(); d = vendor_dir(v); raw = d / 'docs' / 'raw'
xlsx = Path(sys.argv[2]) if len(sys.argv) > 2 else next((p for p in archivos_raw(raw) if p.suffix.lower() in ('.xlsx', '.xlsm') and es_rubrica_implementacion(p)), None)
if not xlsx: sys.exit(f"No encontré la rúbrica de implementación en vendors/{v}/docs/raw (xlsx con hoja «Criterios Estrat Implementación»).")
print('Rúbrica:', xlsx.relative_to(raw))
wb = load_workbook(xlsx, read_only=True, data_only=True)
lim = lambda s, n=4000: re.sub(r'\s*\n\s*', ' / ', str(s).strip())[:n] if s not in (None, '') else ''

def hoja(nombre):
    n = next((s for s in wb.sheetnames if s.strip().lower() == nombre.lower()), None)
    return list(wb[n].iter_rows(values_only=True)) if n else None

def cabecera(rows, clave='Item'):
    for i, r in enumerate(rows[:20]):
        if r and clave in [str(c).strip() for c in r if c]: return i, [str(c).strip() if c else '' for c in r]
    return None, None

def col(hdr, *prefijos):
    for p in prefijos:
        for i, c in enumerate(hdr):
            if c.lower().replace('  ', ' ').startswith(p.lower()): return i
    return None

def escalas_laterales(rows, h):
    """Tablas «Categoría: X (F7, F8)» / «(I19) Categoría: …» + filas Puntaje|Descripción, en cualquier columna."""
    out = {}  # clave (código normalizado o nombre de categoría) → texto de escala
    for ci in range(max(len(r) for r in rows if r)):
        ri = h
        while ri < len(rows):
            cell = rows[ri][ci] if ci < len(rows[ri]) else None
            s = str(cell).strip() if cell not in (None, '') else ''
            if s.startswith('Categoría:') or re.match(r'^\(?[FI]-?\d', s) and 'Categor' in s:
                codigos = re.findall(r'\b([FI])\s*-?\s*(\d{1,2})\b', s)
                cat = re.sub(r'^\(.*?\)\s*', '', s).replace('Categoría:', '').strip()
                cat = re.sub(r'\s*\(.*?\)\s*$', '', cat).strip()
                lineas = []; rj = ri + 1
                while rj < len(rows) and rj < ri + 12:
                    r2 = rows[rj]; a = r2[ci] if ci < len(r2) else None; b = r2[ci + 1] if ci + 1 < len(r2) else None
                    if a in (None, '') and b in (None, ''): break
                    sa = str(a).strip() if a not in (None, '') else ''
                    if sa in ('Puntaje',): rj += 1; continue
                    if re.fullmatch(r'\d+(\.\d+)?', sa): lineas.append(f"{sa} {lim(b, 200)}".strip())
                    rj += 1
                texto = ' / '.join(lineas)
                if texto:
                    for L, n in codigos: out[f"{L}-{int(n):02d}"] = texto
                    if cat: out[cat.lower()] = texto
                ri = rj
            else: ri += 1
    return out

# ---- respuestas del vendor en docs/txt (Anexo de perfiles) ----
resp = {}
for t in (d / 'docs' / 'txt').glob('*.txt'):
    hoja_act = None
    for l in t.read_text(encoding='utf-8', errors='replace').splitlines():
        if l.startswith('===== SHEET'): hoja_act = l; continue
        m = re.match(r'^f\.\d+:\s*\|?\s*(?:[^|]*\|\s*)?(F-\d{2}|I\d-\d{2})\s*\|(.*)$', l)
        if m:
            celdas = [c.strip() for c in m.group(2).split('|')]
            if celdas and celdas[-1]: resp[(m.group(1), t.stem)] = celdas[-1]
def respuesta(codigo):
    hits = [(k[1], val) for k, val in resp.items() if k[0] == codigo]
    return hits

out = [f"# Rúbrica de implementación — {v}", f"# Fuente: {xlsx.relative_to(raw)} | Escala 0–5 por criterio (texto exacto en «escala») | Las sesiones de referenciación no se incluyen (son entrevistas)", ""]
total = 0
# ---- bloque EI: estrategia ----
rows = hoja('Criterios Estrat Implementación')
if rows:
    h, hdr = cabecera(rows)
    iI, iCat, iCri, iCon, iPes, iMax, iEsc = col(hdr, 'Item'), col(hdr, 'Categoría'), col(hdr, 'Criterio'), col(hdr, 'Consideraciones'), col(hdr, 'Peso Relativo'), col(hdr, 'Puntaje Máximo'), col(hdr, 'Escala')
    out.append("## BLOQUE EI — Estrategia de implementación (se evalúa contra la propuesta y anexos del vendor)")
    for r in rows[h + 1:]:
        if not r or r[iI] in (None, '') or not re.fullmatch(r'\d+', str(r[iI]).strip()): continue
        total += 1
        out += [f"\n[EI-{int(r[iI]):02d}] bloque=Estrategia categoria={lim(r[iCat])} peso={lim(r[iPes])} max={lim(r[iMax])}",
                f"CRITERIO: {lim(r[iCri])}", f"CONSIDERACIONES: {lim(r[iCon])}", f"ESCALA: {lim(r[iEsc])}", "CALIF_BANCO: —", "RESPUESTA_VENDOR: (ver propuesta; no hay celda específica)"]
# ---- bloques F e I ----
for nombre, pref, bloque in (('Criterios Fabricante', 'F', 'Fabricante'), ('Criterios Implementador', 'I', 'Implementador')):
    rows = hoja(nombre)
    if not rows: continue
    h, hdr = cabecera(rows); esc = escalas_laterales(rows, h)
    iI, iCat, iCri, iPes, iMax = col(hdr, 'Item'), col(hdr, 'Categoría'), col(hdr, 'Criterio'), col(hdr, 'Peso Relativo'), col(hdr, 'Puntaje Máximo')
    iResp = col(hdr, 'Responsable'); iCal = [i for i, c in enumerate(hdr) if c.lower().startswith('calificacion')]; iPun = [i for i, c in enumerate(hdr) if c.lower() == 'puntaje']
    out.append(f"\n## BLOQUE {pref} — Perfil del {bloque.lower()} (respuestas del vendor en su anexo de perfiles; se verifican contra los documentos)")
    for r in rows[h + 1:]:
        if not r or r[iI] in (None, '') or not re.match(rf'^{pref}-\d{{2}}$', str(r[iI]).strip()): continue
        total += 1; cod = str(r[iI]).strip()
        peso = lim(r[iPes]); informativo = peso.lower() in ('no aplica', '') or (iResp is not None and str(r[iResp]).strip().lower() == 'informativo')
        escala = esc.get(cod) or esc.get(lim(r[iCat]).lower(), '')
        calif = ' / '.join(f"impl{k+1}={lim(r[i])}" for k, i in enumerate(iCal) if i < len(r) and r[i] not in (None, '')) or '—'
        out += [f"\n[{cod}] bloque={bloque} categoria={lim(r[iCat])} peso={'informativo' if informativo else peso} max={lim(r[iMax]) if not informativo else '—'}" + (f" responsable={lim(r[iResp])}" if iResp is not None and r[iResp] else ''),
                f"CRITERIO: {lim(r[iCri])}", f"ESCALA: {escala or '(sin escala específica: usar criterio experto 0–5 y decirlo)'}", f"CALIF_BANCO: {calif}"]
        # respuestas: F-xx directo; I-xx → I1-xx, I2-xx
        codigos = [cod] if pref == 'F' else [f"I1-{cod[2:]}", f"I2-{cod[2:]}"]
        hubo = False
        for c in codigos:
            for doc, val in respuesta(c):
                hubo = True; out.append(f"RESPUESTA_VENDOR ({c} en {doc}): {lim(val, 1500)}")
        if not hubo: out.append("RESPUESTA_VENDOR: (no encontrada en docs/txt)")
out.append(f"\n# Total: {total} criterios")
(d / 'matriz' / 'IMPL.txt').write_text('\n'.join(out) + '\n', encoding='utf-8')
print(f"matriz/IMPL.txt: {total} criterios")
