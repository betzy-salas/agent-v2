#!/usr/bin/env python3
"""Resumen ejecutivo de la matriz técnica de un vendor en formato de presentación (PDF 16:9 con el diseño del deck corporativo y
PowerPoint), a partir de vendors/<vendor>/informes/12_resumen_matriz.md (script 18 + agente resumen-matriz). No calcula ni redacta
nada nuevo: copia los puntajes, niveles y fundamentos ya escritos.

Diapositivas: 1) resultado por hoja; 2) una por hoja: tabla de categorías (puntaje, nivel, Indispensables con 0 o 1) y el fundamento
de la hoja; 3) con --completo, además el fundamento de cada categoría (dos por diapositiva).

Uso:  python scripts/19_resumen_matriz_pptx.py <vendor> [--completo] [--base carpeta] [--logo ruta]
   → <base>/RESUMEN_MATRIZ_<vendor>.html, .pdf y .pptx   (base por defecto: vendors/<vendor>/informes; logo: presentacion_t5*/assets/logo.png si existe)
Requiere Chrome (PDF) y python-pptx (PPTX); si falta alguno, genera lo que pueda y lo avisa.
"""
import datetime, html, importlib.util, re, sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent; RAIZ = AQUI.parent
def _cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, AQUI / archivo); mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod
p10 = _cargar('p10', '10_presentacion_t5.py'); p13 = _cargar('p13', '13_laminas_correo.py'); m = p10.m
esc = html.escape
MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
COLOR_NIVEL = {'cumple plenamente': '1BAF7A', 'cumple': '2A78D6', 'parcial': 'EDA100', 'no cumple': 'D64545'}
def color_nivel(n): return COLOR_NIVEL.get(m._norm(n), '8A8F98')
SIN = '(fundamento no redactado todavía: ejecuta el agente resumen-matriz)'
CSS = p13.CSS + """
.rm table.t{font-size:8.6pt} .rm table.t td{padding:4pt 7pt;vertical-align:top;line-height:1.3} .rm table.t th{padding:5pt 7pt;font-size:8.4pt}
.rm .niv{display:inline-block;padding:1pt 7pt;border-radius:3pt;font-weight:700;color:#fff;font-size:8pt;white-space:nowrap}
.rm .pt{font-family:"JetBrains Mono",Consolas,monospace;font-weight:800;color:#00205b}
.rm .dos{display:flex;gap:16pt;align-items:flex-start} .rm .dos>div{min-width:0} .rm .fund{font-size:9.6pt;line-height:1.45;color:#171f2a;margin:0}
.rm .fund.chico{font-size:8.8pt} .rm h4{margin:0 0 4pt;font-size:11pt;color:#00205b} .rm .cat{margin-bottom:12pt;padding-bottom:10pt;border-bottom:1px solid #e3e6ec} .rm .cat:last-child{border:0}
.rm .sub{font-size:8.4pt;color:#5f6b7a;margin:0 0 6pt}
"""

def leer(vendor):
    f = m.VENDORS / vendor / 'informes' / '12_resumen_matriz.md'
    if not f.exists(): sys.exit(f'Falta {f}: ejecuta primero scripts/18_resumen_matriz.py {vendor} (y el agente resumen-matriz).')
    md = f.read_text(encoding='utf-8')
    cab = md.split('\n## ', 1)[0]
    matriz = (re.search(r'- Matriz:\s*(.+)', cab) or [None, ''])[1].strip()
    hdr, rows = m.first_table(m.find_section(m.sections(md), 'Resultado por hoja'))
    hojas = []
    for bloque in re.split(r'\n(?=## Hoja )', md)[1:]:
        if bloque.startswith('## Nota'): continue
        tit = bloque.split('\n', 1)[0][3:].strip(); nombre = tit.split(' · ')[0].replace('Hoja ', '').strip(); dim = tit.split(' · ', 1)[1].strip() if ' · ' in tit else ''
        pj = re.search(r'\*\*Puntaje [^:]+: ([^/]+) / 100 · ([^*]+)\*\*', bloque); punt, niv = (pj.group(1).strip(), pj.group(2).strip()) if pj else ('—', '—')
        pre = bloque.split('### Categorías', 1)[0]
        parrafos = [p.strip() for p in pre.split('\n\n') if p.strip() and not p.strip().startswith(('## ', '**Puntaje', '<!--'))]
        fund_hoja = parrafos[-1] if parrafos else SIN
        chdr, crows = m.first_table(bloque.split('### Categorías', 1)[1] if '### Categorías' in bloque else '')
        cats = []
        for cb in re.split(r'\n(?=#### )', bloque)[1:]:
            t = cb.split('\n', 1)[0][5:]; mm = re.match(r'(.+?) — ([^/]+) / 100 · (.+)$', t.strip())
            cnombre, cp, cn = (mm.group(1), mm.group(2).strip(), mm.group(3).strip()) if mm else (t, '—', '—')
            cuerpo = cb.split('\n', 1)[1] if '\n' in cb else ''
            pars = [p.strip() for p in cuerpo.split('\n\n') if p.strip() and not p.strip().startswith(('|', '<!--', '## ', '### '))]
            cats.append({'nombre': cnombre.strip(), 'punt': cp, 'nivel': cn, 'fund': pars[0] if pars else SIN, 'fila': next((r for r in crows if r and r[0].strip() == cnombre.strip()), None)})
        hojas.append({'hoja': nombre, 'dim': dim, 'punt': punt, 'nivel': niv, 'fund': fund_hoja, 'chdr': chdr, 'cats': cats})
    return {'vendor': vendor, 'nombre': m.nombre(vendor), 'matriz': matriz, 'hdr': hdr, 'rows': rows, 'hojas': hojas}

def ix(hdr, *pref):
    for p in pref:
        for i, h in enumerate(hdr):
            if m._norm(h).startswith(m._norm(p)): return i
    return None

# ------------------------------------------------------------------ HTML / PDF
def construir_html(d, logo, completo):
    t = datetime.date.today(); mes = f'{MESES[t.month - 1].capitalize()} {t.year}'
    def bar(titulo, derecha=''): return f'<div class="bar"><h2>{esc(titulo)}</h2><span class="r">{esc(derecha)}</span></div>'
    def pie(fuente): return f'<div class="pie"><span>{esc(fuente)}</span><span class="c">Datos elaborados por BCP para uso Interno</span>' + (f'<img src="{logo}" alt="">' if logo else '') + '</div>'
    def niv(n): return f'<span class="niv" style="background:#{color_nivel(n)}">{esc(n)}</span>'
    fuente = f'Fuente: matriz técnica del Banco ({d["matriz"]}), calificaciones tal cual; resumen por categorías (informes/12_resumen_matriz.md)'
    S = []
    # 1) resultado por hoja
    H = d['hdr']; iH, iD, iI, iP, iN, iS, iInc = ix(H, 'hoja'), ix(H, 'dimensi'), ix(H, 'items', 'ítems'), ix(H, 'puntaje'), ix(H, 'nivel'), ix(H, 'sin calificar'), ix(H, 'incoheren')
    trs = ''.join(f'<tr><td class="k">{esc(r[iH])}</td><td>{esc(r[iD]) if iD is not None else ""}</td><td class="c">{esc(r[iI])}</td><td class="c pt">{esc(r[iP])}</td><td>{niv(r[iN])}</td><td class="c">{esc(r[iS])}</td><td class="c">{esc(r[iInc]) if iInc is not None else ""}</td></tr>' for r in d['rows'])
    nombre = d['nombre']
    S.append(f'<section class="s rm">{bar("Resumen ejecutivo de la matriz técnica · " + nombre, "Calificaciones del Banco tal cual · " + mes)}<div class="body">'
             f'<p class="sub">Puntaje = media de las calificaciones (0–4) sobre 4, ponderada por criticidad (Indispensable/Killer 3, Deseable/No killer 1). Nivel: ≥ 85 Cumple plenamente · 70–84 Cumple · 50–69 Parcial · &lt; 50 No cumple. Sin recalificación.</p>'
             f'<table class="t"><thead><tr><th>Hoja</th><th>Dimensión</th><th class="c">Ítems</th><th class="c">Puntaje</th><th>Nivel</th><th class="c">Sin calificar</th><th class="c">Incoherencias Eval/VA</th></tr></thead><tbody>{trs}</tbody></table></div>{pie(fuente)}</section>')
    # 2) una por hoja
    for h in d['hojas']:
        C = h['chdr']; iC, iIt, iPc, iNc, iB = ix(C, 'categor'), ix(C, 'items', 'ítems'), ix(C, 'puntaje'), ix(C, 'nivel'), ix(C, 'indispens')
        filas = [c['fila'] for c in h['cats'] if c['fila']]
        chico = len(filas) > 11
        trs = ''.join(f'<tr><td class="k">{esc(r[iC])}</td><td class="c">{esc(r[iIt])}</td><td class="c pt">{esc(r[iPc])}</td><td>{niv(r[iNc])}</td><td>{esc(r[iB]) if iB is not None else ""}</td></tr>' for r in filas)
        tabla = f'<table class="t" style="font-size:{7.2 if chico else 8.4}pt"><thead><tr><th>Categoría</th><th class="c">Ítems</th><th class="c">Puntaje</th><th>Nivel</th><th>Indisp. con 0 o 1</th></tr></thead><tbody>{trs}</tbody></table>'
        hj, dim, sub = h['hoja'], h['dim'] or nombre, 'Puntaje ' + h['punt'] + ' / 100 · ' + h['nivel']
        S.append(f'<section class="s rm">{bar("Hoja " + hj + " · " + dim, sub)}<div class="body"><div class="dos">'
                 f'<div style="flex:0 0 52%">{tabla}</div><div style="flex:1"><h4>En qué se fundamenta el resultado de la hoja</h4><p class="fund{" chico" if len(h["fund"]) > 900 else ""}">{esc(h["fund"])}</p></div></div></div>{pie(fuente)}</section>')
        if completo:
            cats = h['cats']
            for k in range(0, len(cats), 2):
                par = cats[k:k + 2]
                cuerpo = ''.join(f'<div class="cat"><h4>{esc(c["nombre"])} <span class="pt" style="margin-left:6pt">{esc(c["punt"])}</span> {niv(c["nivel"])}</h4><p class="fund{" chico" if len(c["fund"]) > 800 else ""}">{esc(c["fund"])}</p></div>' for c in par)
                S.append(f'<section class="s rm">{bar("Hoja " + hj + " · fundamento por categoría", dim + " · " + str(k // 2 + 1) + " de " + str((len(cats) + 1) // 2))}<div class="body">{cuerpo}</div>{pie(fuente)}</section>')
    return f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Resumen matriz · {esc(d["nombre"])}</title><style>{CSS}</style></head><body>' + ''.join(S) + '</body></html>'

# ------------------------------------------------------------------ PPTX
def construir_pptx(d, logo_path, completo, salida):
    try:
        p12 = _cargar('p12', '12_presentacion_t5_pptx.py'); p14 = _cargar('p14', '14_laminas_correo_pptx.py')
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.text import MSO_ANCHOR
    except (ImportError, SystemExit):
        print('AVISO: sin python-pptx no se genera el .pptx (pip install python-pptx)'); return 0
    NAVY_T, GREY = p12.NAVY_T, p12.GREY
    t = datetime.date.today(); mes = f'{MESES[t.month - 1].capitalize()} {t.year}'
    fuente = f'Fuente: matriz técnica del Banco ({d["matriz"]}), calificaciones tal cual; resumen por categorías (informes/12_resumen_matriz.md)'
    prs = Presentation(); prs.slide_width = Inches(p12.W); prs.slide_height = Inches(p12.H)
    def nueva(): return prs.slides.add_slide(prs.slide_layouts[6])
    def niv_run(n): return (n, {'bold': True, 'color': color_nivel(n)})
    s = nueva(); p12.barra(s, f'Resumen ejecutivo de la matriz técnica · {d["nombre"]}', 'Calificaciones del Banco tal cual · ' + mes)
    p12.texto(s, 0.42, 0.72, 12.5, 0.4, [[('Puntaje = media de las calificaciones (0–4) sobre 4, ponderada por criticidad (Indispensable/Killer 3, Deseable/No killer 1). Nivel: ≥ 85 Cumple plenamente · 70–84 Cumple · 50–69 Parcial · < 50 No cumple. Sin recalificación.', {'color': GREY, 'size': 8.5})]], ajustar=False)
    H = d['hdr']; iH, iD, iI, iP, iN, iS, iInc = ix(H, 'hoja'), ix(H, 'dimensi'), ix(H, 'items', 'ítems'), ix(H, 'puntaje'), ix(H, 'nivel'), ix(H, 'sin calificar'), ix(H, 'incoheren')
    cuerpo = [[r[iH], r[iD] if iD is not None else '', r[iI], [(r[iP], {'bold': True, 'font': p12.MONO, 'color': NAVY_T})], [niv_run(r[iN])], r[iS], r[iInc] if iInc is not None else ''] for r in d['rows']]
    p12.tabla(s, 0.42, 1.2, 12.5, [0.8, 3.6, 0.8, 1.0, 1.9, 1.2, 1.6], ['Hoja', 'Dimensión', 'Ítems', 'Puntaje', 'Nivel', 'Sin calificar', 'Incoherencias Eval/VA'], cuerpo, size=9, head_size=9, center_cols=(2, 3, 5, 6))
    p14.pie(s, fuente, logo_path)
    for h in d['hojas']:
        C = h['chdr']; iC, iIt, iPc, iNc, iB = ix(C, 'categor'), ix(C, 'items', 'ítems'), ix(C, 'puntaje'), ix(C, 'nivel'), ix(C, 'indispens')
        filas = [c['fila'] for c in h['cats'] if c['fila']]; n = len(filas); sz = 8 if n <= 11 else 7 if n <= 16 else 6
        s = nueva(); p12.barra(s, f'Hoja {h["hoja"]} · {h["dim"] or d["nombre"]}', f'Puntaje {h["punt"]} / 100 · {h["nivel"]}')
        cuerpo = [[r[iC], r[iIt], [(r[iPc], {'bold': True, 'font': p12.MONO, 'color': NAVY_T})], [niv_run(r[iNc])], r[iB] if iB is not None else ''] for r in filas]
        p12.tabla(s, 0.42, 0.8, 6.5, [2.6, 0.6, 0.8, 1.4, 1.1], ['Categoría', 'Ítems', 'Puntaje', 'Nivel', 'Indisp. 0/1'], cuerpo, size=sz, head_size=sz, center_cols=(1, 2), row_h=0.26 if n > 16 else None)
        p14.h3(s, 7.2, 0.8, 5.7, 'En qué se fundamenta el resultado de la hoja', size=12)
        p12.texto(s, 7.2, 1.2, 5.7, 5.7, [[(h['fund'], {'size': 9})]], size=9, min_size=6.5)
        p14.pie(s, fuente, logo_path)
        if completo:
            cats = h['cats']
            for k in range(0, len(cats), 2):
                s = nueva(); p12.barra(s, f'Hoja {h["hoja"]} · fundamento por categoría', f'{h["dim"] or d["nombre"]} · {k // 2 + 1} de {(len(cats) + 1) // 2}')
                y = 0.8
                for c in cats[k:k + 2]:
                    p12.texto(s, 0.42, y, 12.5, 0.32, [[(c['nombre'], {'bold': True, 'color': NAVY_T, 'size': 12}), (f'   {c["punt"]}  ', {'bold': True, 'font': p12.MONO, 'color': NAVY_T, 'size': 11}), niv_run(c['nivel'])]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
                    p12.texto(s, 0.42, y + 0.36, 12.5, 2.6, [[(c['fund'], {'size': 9.5})]], size=9.5, min_size=7)
                    y += 3.1
                p14.pie(s, fuente, logo_path)
    prs.save(salida); return 1

def main():
    args = sys.argv[1:]
    if not args or args[0].startswith('-'): sys.exit('Uso: python scripts/19_resumen_matriz_pptx.py <vendor> [--completo] [--base carpeta] [--logo ruta]')
    vendor = args[0]; completo = '--completo' in args
    base = Path(args[args.index('--base') + 1]) if '--base' in args else m.VENDORS / vendor / 'informes'
    if not base.is_absolute(): base = RAIZ / base
    base.mkdir(parents=True, exist_ok=True)
    logo_path = Path(args[args.index('--logo') + 1]) if '--logo' in args else next((p for p in sorted(RAIZ.glob('presentacion_t5*/assets/logo.png'))), None)
    d = leer(vendor)
    out = base / f'RESUMEN_MATRIZ_{vendor}.html'; out.write_text(construir_html(d, p10._data_uri(logo_path) if logo_path and logo_path.exists() else '', completo), encoding='utf-8'); print(f'HTML: {out}')
    pdf = out.with_suffix('.pdf'); m.a_pdf(out, pdf)
    try:
        import fitz; n = fitz.open(pdf).page_count; esperadas = 1 + len(d['hojas']) + (sum((len(h['cats']) + 1) // 2 for h in d['hojas']) if completo else 0)
        print(f'PDF: {pdf} · {n} páginas' + ('' if n == esperadas else f' (AVISO: se esperaban {esperadas}; alguna diapositiva se desbordó)'))
    except ImportError: pass
    if construir_pptx(d, logo_path, completo, out.with_suffix('.pptx')): print(f'PPTX: {out.with_suffix(".pptx")}')

if __name__ == '__main__':
    main()
