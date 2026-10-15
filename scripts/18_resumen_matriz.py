#!/usr/bin/env python3
"""Resumen por categorías de la matriz técnica del Banco, hoja por hoja, con las calificaciones TAL CUAL están (sin recalificar).

Sirve para cualquier matriz con la estructura del Banco, tenga las hojas que tenga (T1..T8 o un documento nuevo con otras hojas):
una hoja cuenta como «hoja de matriz» si tiene una fila de cabecera con la columna N° y la columna VA (cabeceras en scripts/_comun.py).
Por cada hoja detectada:
  - escribe vendors/<vendor>/matriz/<hoja>.txt en el formato del pipeline ([E<n>] cat= crit= va= eval= / Q / PROVEEDOR / SUSTENTO),
  - calcula por categoría: ítems, calificados, puntaje 0–100 (media de VA sobre 4 ponderada por criticidad: Indispensable/Killer 3,
    Deseable/No killer 1, como el consolidado del Banco), nivel oficial (CLAUDE.md: ≥85 Cumple plenamente · 70–84 Cumple · 50–69 Parcial
    · <50 No cumple), reparto de calificaciones 0..4, ítems sin calificar e incoherencias Eval/VA,
  - lista los ítems de cada categoría (N°, criticidad, calificación, pregunta y sustento del evaluador recortados),
  - deja marcadores <!-- resumen-hoja: <hoja> --> y <!-- resumen: <hoja> | <categoría> --> que completa el agente resumen-matriz.
Salida: vendors/<vendor>/informes/12_resumen_matriz.md (o --salida).

Uso:  python scripts/18_resumen_matriz.py <vendor> [ruta_al_xlsx] [--hojas A,B,C] [--salida archivo.md] [--limpiar]
      --limpiar: aparta a matriz/_anterior/ los <hoja>.txt que no correspondan a hojas de esta matriz (p. ej. los T1..T8 de la matriz anterior); IMPL.txt no se toca.
      Sin ruta, busca la matriz en matrices/*<vendor>*.xlsx (la más reciente), luego en vendors/<vendor>/docs/raw y por último en vendors/<vendor>/.
"""
import datetime, re, sys
from collections import OrderedDict
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _comun import arg_vendor, vendor_dir, archivos_raw, es_matriz_evaluacion, COLUMNAS as K, requiere, RAIZ
requiere('openpyxl')
import openpyxl

ESCALA = {0: 'No Cubierto', 1: 'Parcialmente Cubierto', 2: 'Cubierto: Modificación Mayor', 3: 'Cubierto: Modificación Menor', 4: 'Completamente Cubierto'}
PESO_CRIT = {'indispensable': 3, 'killer': 3, 'deseable': 1, 'no killer': 1}
def peso(it): return PESO_CRIT.get((it.get('crit') or '').strip().lower(), 1)
def puntaje(its):
    vas = [(peso(i), i['va']) for i in its if i['va'] is not None]
    return round(sum(p * v for p, v in vas) / (4 * sum(p for p, _ in vas)) * 100, 1) if vas else None
def nivel(p):
    if p is None: return '—'
    return 'Cumple plenamente' if p >= 85 else 'Cumple' if p >= 70 else 'Parcial' if p >= 50 else 'No cumple'
def norm(s): return re.sub(r'\s+', ' ', str(s or '')).strip()
def corta(t, n): t = norm(t); return t if len(t) <= n else t[:n - 1].rstrip() + '…'
def celda(t): return norm(t).replace('|', '/')

def en_matrices(v):
    """matrices/<algo con el vendor>.xlsx: carpeta del vendor o su nombre oficial en el nombre del archivo; si hay varios, el más reciente."""
    carpeta = RAIZ / 'matrices'
    if not carpeta.exists(): return None
    import importlib.util as _iu
    spec = _iu.spec_from_file_location('m07', Path(__file__).resolve().parent / '07_informe_implementacion.py'); m07 = _iu.module_from_spec(spec); spec.loader.exec_module(m07)
    claves = {re.sub(r'[^a-z0-9]', '', v.lower()), re.sub(r'[^a-z0-9]', '', m07.nombre(v).lower())}
    cands = [p for p in carpeta.glob('*.xls*') if not p.name.startswith('~$') and any(k and k in re.sub(r'[^a-z0-9]', '', p.stem.lower()) for k in claves)]
    if not cands: return None
    cands.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    if len(cands) > 1: print(f'⚠ Varias matrices para {v} en matrices/: uso la más reciente, {cands[0].name}')
    return cands[0]

def _es_numero(c):
    t = re.sub(r'[\s.]', '', str(c)).lower()
    return t in ('n°', 'nº', 'no', 'n', 'num', 'número', 'numero', 'item', 'ítem', 'id') or t == re.sub(r'[\s.]', '', K['numero']).lower()
def _es_va(c): return str(c).strip().upper() in (K['va'].upper(), 'V.A.', 'V.A', 'VA.', 'VALOR', 'CALIF', 'CALIFICACION', 'CALIFICACIÓN')
def hojas_matriz(wb, informar=False):
    """Hojas con una fila de cabecera que tenga la columna N° (o Nº, No., N) y la columna VA (en las primeras 40 filas), en el orden del
    libro; devuelve [(nombre, filas, índice de cabecera)]. Con informar=True imprime por qué se descarta cada hoja no reconocida."""
    out = []
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True)); hallado = None; tiene_n = tiene_va = False
        for h, r in enumerate(rows[:40]):
            cells = [c for c in r if c is not None and str(c).strip()]
            n_ok = any(_es_numero(c) for c in cells); va_ok = any(_es_va(c) for c in cells)
            tiene_n |= n_ok; tiene_va |= va_ok
            if n_ok and va_ok: hallado = h; break
        if hallado is not None: out.append((name, rows, hallado))
        elif informar:
            motivo = 'sin columna VA' if tiene_n and not tiene_va else 'sin columna N°' if tiene_va and not tiene_n else 'sin cabecera N° + VA en las primeras 40 filas'
            print(f'  hoja «{name}»: no es hoja de matriz ({motivo})')
    return out

def leer_hoja(rows, h):
    hdr = [str(c).strip() if c else '' for c in rows[h]]
    col = lambda name: next((i for i, c in enumerate(hdr) if c.startswith(name)), None)
    exact = lambda name: next((i for i, c in enumerate(hdr) if c == name), None)
    iN, iCat, iQ, iCrit, iCrit2 = col(K['numero']), col(K['categoria']), col(K['pregunta']), col(K['criticidad']), col(K['criticidad_nueva'])
    iEst, iSus, iEval, iEvdor, iDim = col(K['estado']), col(K['sustento']), exact(K['eval_texto']), col(K['evaluador']), col('Dimensi')
    iVA = next((i for i, c in enumerate(hdr) if _es_va(c)), None)
    if iN is None: iN = next((i for i, c in enumerate(hdr) if _es_numero(c)), None)
    iResp = [i for i, c in enumerate(hdr) if c.startswith(K['respuesta'])]
    g = lambda r, i: (r[i] if i is not None and i < len(r) else None)
    items = []; dimension = ''
    for r in rows[h + 1:]:
        if not g(r, iN) or not str(g(r, iN)).startswith('E'): continue
        if iEst is not None and str(g(r, iEst)).strip() == 'Eliminar': continue
        if g(r, iQ) and str(g(r, iQ)).strip().lower().startswith(('se elimin', 'se elemin')): continue
        crit = (g(r, iCrit2) if iCrit2 is not None and g(r, iCrit2) else g(r, iCrit)) or ''
        resp = ' | '.join(str(r[i]) for i in iResp if i < len(r) and r[i] and str(r[i]).strip() not in ('Sin cambios.', 'None', 'nan'))
        va = g(r, iVA); va = int(va) if isinstance(va, (int, float)) else None
        ev = g(r, iEval); ev = norm(ev) if ev not in (None, '') else None
        if not dimension and g(r, iDim): dimension = norm(g(r, iDim))
        items.append({'n': str(g(r, iN)).strip(), 'cat': norm(g(r, iCat)) or '(sin categoría)', 'crit': norm(crit), 'q': norm(g(r, iQ)), 'resp': resp,
                      'va': va, 'eval': ev, 'evaluador': norm(g(r, iEvdor)) or None, 'sus': norm(g(r, iSus))})
    return items, dimension

def incoherente(it):
    """Eval textual que no corresponde al VA según la escala (p. ej. eval=Cubierto: Modificación Menor y va=4)."""
    if it['va'] is None or not it['eval'] or it['eval'].lower().startswith('seleccionar'): return False
    esperado = ESCALA.get(it['va'], ''); return re.sub(r'\W', '', it['eval'].lower()) != re.sub(r'\W', '', esperado.lower())

def escribir_txt(d, hoja, items):
    out = d / 'matriz'; out.mkdir(exist_ok=True); f = out / f'{hoja}.txt'
    with open(f, 'w', encoding='utf-8') as o:
        o.write(f'# Hoja {hoja}\n')
        for it in items:
            o.write(f"\n[{it['n']}] cat={it['cat']} crit={it['crit']} va={it['va']} eval={it['eval']} evaluador={it['evaluador']}\nQ: {it['q']}\nPROVEEDOR: {it['resp']}\nSUSTENTO: {it['sus']}\n")
    return f

def main():
    v = arg_vendor(); d = vendor_dir(v); args = sys.argv[2:]
    def opc(n):
        if n in args: i = args.index(n); val = args[i + 1]; del args[i:i + 2]; return val
    hojas_pedidas = opc('--hojas'); salida = opc('--salida'); limpiar = '--limpiar' in args
    if limpiar: args.remove('--limpiar')
    if args and not args[0].startswith('-'): xlsx = Path(args[0])
    else:
        xlsx = en_matrices(v)
    if xlsx is None:
        cands = [p for p in archivos_raw(d / 'docs' / 'raw') + sorted(d.glob('*.xls*')) if p.suffix.lower() in ('.xlsx', '.xlsm') and es_matriz_evaluacion(p)]
        if not cands:   # sin hoja de método ni T1–T8: cualquier xlsx del vendor con alguna hoja de matriz
            for p in archivos_raw(d / 'docs' / 'raw') + sorted(d.glob('*.xls*')):
                if p.suffix.lower() in ('.xlsx', '.xlsm') and hojas_matriz(openpyxl.load_workbook(p, read_only=True, data_only=True)): cands.append(p)
        if not cands: sys.exit(f'No encontré una matriz (xlsx con columnas N° y VA) en matrices/ (nombre con «{v}»), vendors/{v}/docs/raw ni vendors/{v}/. Pásala como segundo argumento.')
        if len(cands) > 1: print('⚠ Varias matrices candidatas, uso la primera:', [c.name for c in cands])
        xlsx = cands[0]
    print('Matriz:', xlsx)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    print('Hojas del libro:', ', '.join(wb.sheetnames))
    hojas = hojas_matriz(wb, informar=True)
    print('Hojas de matriz detectadas:', ', '.join(h for h, _, _ in hojas) or 'ninguna')
    if hojas_pedidas: pedidas = [h.strip() for h in hojas_pedidas.split(',')]; hojas = [h for h in hojas if h[0] in pedidas]
    if not hojas: sys.exit('Ninguna hoja del Excel tiene cabecera con N° y VA.')
    hoy = datetime.date.today().strftime('%d-%m-%Y'); rel = xlsx.relative_to(RAIZ) if str(xlsx).startswith(str(RAIZ)) else xlsx
    L = [f'# 12. Resumen de la matriz técnica por categorías — {v}', '',
         f'- Vendor: {v}', f'- Matriz: {rel} (hojas detectadas: {", ".join(h for h, _, _ in hojas)})', f'- Fecha: {hoy}',
         '- Calificaciones: las del Banco (columna VA, 0–4) tal como están en la matriz; este documento no recalifica ni propone calificaciones.',
         '- Puntaje de hoja y de categoría: media de las calificaciones (0–4) sobre 4, ponderada por criticidad (Indispensable/Killer 3, Deseable/No killer 1), en 0–100; los ítems sin calificar no entran.',
         '- Nivel: escala oficial de CLAUDE.md («Niveles de cumplimiento»): ≥ 85 Cumple plenamente · 70–84 Cumple · 50–69 Parcial · < 50 No cumple.',
         '- Fundamento por hoja y por categoría (en qué se basa el resultado: qué pide el Banco, qué ofrece el proveedor, qué le faltó según el evaluador): lo escribe el agente resumen-matriz a partir de la pregunta, la respuesta del proveedor y el sustento del evaluador (matriz/<hoja>.txt), sin añadir hechos ni repetir conteos.', '',
         '## Resultado por hoja', '', '| Hoja | Dimensión | Ítems | Calificados | Puntaje | Nivel | 4 | 3 | 2 | 1 | 0 | Sin calificar | Incoherencias Eval/VA |', '|---|---|---|---|---|---|---|---|---|---|---|---|---|']
    detalle = []
    for hoja, rows, h in hojas:
        items, dim = leer_hoja(rows, h)
        if not items: print(hoja, 'sin ítems; se omite'); continue
        f = escribir_txt(d, hoja, items); print(hoja, 'ítems:', len(items), '→', f.relative_to(RAIZ))
        dist = {k: sum(1 for i in items if i['va'] == k) for k in range(5)}; sc = sum(1 for i in items if i['va'] is None); inc = [i['n'] for i in items if incoherente(i)]
        p = puntaje(items)
        L.append(f'| {hoja} | {celda(dim)} | {len(items)} | {len(items) - sc} | {p if p is not None else "—"} | {nivel(p)} | {dist[4]} | {dist[3]} | {dist[2]} | {dist[1]} | {dist[0]} | {sc} | {len(inc)} |')
        cats = OrderedDict()
        for it in items: cats.setdefault(it['cat'], []).append(it)
        D = [f'## Hoja {hoja}' + (f' · {dim}' if dim else ''), '', f'**Puntaje {hoja}: {p if p is not None else "—"} / 100 · {nivel(p)}** · {len(items)} ítems, {len(items) - sc} calificados' + (f' · incoherencias Eval/VA: {", ".join(inc)}' if inc else '') + '.', '',
             f'<!-- resumen-hoja: {hoja} -->', '', '### Categorías', '', '| Categoría | Ítems | Puntaje | Nivel | 4 | 3 | 2 | 1 | 0 | Sin calif. | Indispensables con 0 o 1 |', '|---|---|---|---|---|---|---|---|---|---|---|']
        for c, its in cats.items():
            pc = puntaje(its); dc = {k: sum(1 for i in its if i['va'] == k) for k in range(5)}; scc = sum(1 for i in its if i['va'] is None)
            bajos = [i['n'] for i in its if i['va'] in (0, 1) and peso(i) == 3]
            D.append(f'| {celda(c)} | {len(its)} | {pc if pc is not None else "—"} | {nivel(pc)} | {dc[4]} | {dc[3]} | {dc[2]} | {dc[1]} | {dc[0]} | {scc} | {", ".join(bajos) or "—"} |')
        D += ['', '### Fundamento por categoría', '']
        for c, its in cats.items():
            pc = puntaje(its)
            D += [f'#### {c} — {pc if pc is not None else "—"} / 100 · {nivel(pc)}', '', f'<!-- resumen: {hoja} | {c} -->', '',
                  '| N° | Criticidad | Calificación del Banco | Pregunta | Sustento del evaluador |', '|---|---|---|---|---|']
            for it in its:
                cal = f'{ESCALA[it["va"]]} ({it["va"]})' if it['va'] is not None else 'Sin calificar (—)'
                if incoherente(it): cal = f'{it["eval"]} (VA={it["va"]} ⚠)'
                D.append(f'| {it["n"]} | {celda(it["crit"]) or "—"} | {cal} | {celda(corta(it["q"], 160))} | {celda(corta(it["sus"], 200)) or "—"} |')
            D.append('')
        detalle += D
    if limpiar:   # archivos de hojas de una matriz anterior: se apartan a matriz/_anterior/ (no se borran); IMPL.txt y los .txt de la rúbrica no se tocan
        import shutil
        detectadas = {h for h, _, _ in hojas}; aparte = d / 'matriz' / '_anterior'; movidos = []
        for f in sorted((d / 'matriz').glob('*.txt')):
            if f.stem in detectadas or f.stem.upper().startswith('IMPL'): continue
            aparte.mkdir(exist_ok=True); shutil.move(str(f), str(aparte / f.name)); movidos.append(f.name)
        print('Apartados a matriz/_anterior/:', ', '.join(movidos) if movidos else 'ninguno')
    L += [''] + detalle + ['## Nota', '', 'Generado por scripts/18_resumen_matriz.py desde la matriz del Banco; sin ranking, sin recomendación y sin recalificación. Las verificaciones documentales (verif/) no intervienen.']
    out = Path(salida) if salida else d / 'informes' / '12_resumen_matriz.md'; out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text('\n'.join(L) + '\n', encoding='utf-8'); print('Resumen:', out.relative_to(RAIZ) if str(out).startswith(str(RAIZ)) else out)

if __name__ == '__main__':
    main()
