"""Utilidades compartidas por los scripts del pipeline multi-vendor.

Estructura: vendors/<vendor>/{docs/raw, docs/txt, matriz, verif, informes, INVENTARIO.md, NOMBRES.md}
Todos los scripts reciben el vendor como primer argumento posicional.
"""
import re, sys, pathlib, unicodedata

RAIZ = pathlib.Path(__file__).resolve().parent.parent

# Cabeceras de la matriz de evaluación del Banco (hoja T1..T8 del xlsx). Se comparan por prefijo salvo
# las marcadas exact. Si otra entidad usa otras cabeceras, cámbialas aquí.
COLUMNAS = {
    'numero': 'N°',                              # id del ítem (E1, E2, …)
    'categoria': 'Categoría',
    'pregunta': 'Descripción',
    'criticidad': 'Nivel de criticidad',
    'criticidad_nueva': 'Nuevo Nivel',
    'estado': 'Nuevo / Modificado',              # Eliminar / Omitir Respuesta descartan el ítem
    'respuesta': ('Respuesta del Proveedor', 'Comentario'),   # prefijos de las columnas de respuesta del proveedor (v1, v2…)
    'evaluador': 'Evaluador',                    # nombre del evaluador (prefijo)
    'eval_texto': 'Eval',                        # clasificación textual (exacto)
    'va': 'VA',                                  # calificación numérica 0–4 (exacto)
    'sustento': 'Sustento',                      # justificación escrita del evaluador (prefijo)
    'hoja_metodo': 'Metodo de Evaluacion',       # hoja con la escala
}
VENDORS = RAIZ / 'vendors'
SUBDIRS = ['docs/raw', 'docs/txt', 'matriz', 'verif', 'informes']

def vendor_dir(nombre: str, crear: bool = False) -> pathlib.Path:
    d = VENDORS / nombre
    if not d.exists():
        if not crear:
            disponibles = sorted(p.name for p in VENDORS.iterdir() if p.is_dir()) if VENDORS.exists() else []
            sys.exit(f"No existe vendors/{nombre}. Vendors disponibles: {', '.join(disponibles) or '(ninguno)'}. "
                     f"Crea uno con: python scripts/00_nuevo_vendor.py {nombre}")
        for s in SUBDIRS: (d / s).mkdir(parents=True, exist_ok=True)
    return d

def arg_vendor() -> str:
    if len(sys.argv) < 2 or sys.argv[1].startswith('-'):
        sys.exit(f"Uso: python {pathlib.Path(sys.argv[0]).name} <vendor> [opciones]")
    return sys.argv[1]

def archivos_raw(raw: pathlib.Path) -> list:
    """Todos los documentos de docs/raw, incluidas subcarpetas (se ignoran ocultos y .gitkeep), ordenados."""
    if not raw.exists(): return []
    return sorted(p for p in raw.rglob('*') if p.is_file() and not p.name.startswith('.')
                  and not any(part.startswith('.') for part in p.relative_to(raw).parts))

def slug_de(p: pathlib.Path, raw: pathlib.Path) -> str:
    """Nombre base del txt/toc en docs/txt: el stem si está en la raíz de docs/raw; si está en una
    subcarpeta, «<subcarpetas> - <stem>» para que no choquen archivos homónimos de carpetas distintas."""
    rel = p.relative_to(raw)
    if len(rel.parts) == 1: return p.stem
    return ' - '.join(list(rel.parts[:-1]) + [p.stem])

def nombre_oficial(stem: str) -> str:
    """Nombre legible y estable para citar un documento, derivado del nombre de archivo."""
    s = stem.strip()
    s = re.sub(r'[_]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def escribir_nombres(vdir: pathlib.Path) -> pathlib.Path:
    """Regenera vendors/<v>/NOMBRES.md: tabla slug → nombre oficial → archivo origen → unidad de cita.
    Respeta los nombres oficiales que la persona haya editado a mano (columna 2) si el slug ya existía."""
    txt = vdir / 'docs' / 'txt'; raw = vdir / 'docs' / 'raw'
    destino = vdir / 'NOMBRES.md'
    previos = {}
    if destino.exists():
        for l in destino.read_text(encoding='utf-8').splitlines():
            m = re.match(r'^\|\s*`([^`]+)`\s*\|\s*([^|]+?)\s*\|', l)
            if m: previos[m.group(1)] = m.group(2)
    origen = {slug_de(p, raw): p for p in archivos_raw(raw)}
    filas = []
    for t in sorted(txt.glob('*.txt')):
        stem = t.stem
        src = origen.get(stem)
        ext = src.suffix.lower() if src else ''
        unidad = ('p.<n> (página PDF del marcador PAGE; impresa = n-1)' if ext == '.pdf'
                  else 'hoja <nombre>, f.<fila>' if ext in ('.xlsx', '.xlsm') else 'slide <n>' if ext == '.pptx' else 'p.<n>')
        oficial = previos.get(stem) or nombre_oficial(stem)
        filas.append(f"| `{stem}` | {oficial} | {src.relative_to(raw) if src else '—'} | {unidad} |")
    cuerpo = ["# Nombres oficiales de documento — " + vdir.name, "",
              "Toda cita en verif/ e informes/ usa EXACTAMENTE el texto de la columna «Nombre oficial», seguido de la unidad.",
              "Puedes editar a mano la columna «Nombre oficial» (por ejemplo acortarla); el script la conserva al regenerar.",
              "La tabla la mantiene scripts/01_pdf_a_texto.py y 01b_anexos_a_texto.py.", "",
              "| slug (docs/txt) | Nombre oficial | Archivo en docs/raw | Unidad de cita |", "|---|---|---|---|"] + filas
    destino.write_text('\n'.join(cuerpo) + '\n', encoding='utf-8')
    return destino


def es_matriz_evaluacion(xlsx: pathlib.Path) -> bool:
    """True si el xlsx es la matriz de evaluación del Banco: hoja de método (COLUMNAS['hoja_metodo']) o hojas T1..T8 con la columna VA."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True, data_only=True)
        objetivo = re.sub(r'[^a-z]', '', COLUMNAS['hoja_metodo'].lower())
        if any(re.sub(r'[^a-z]', '', n.lower()) in (objetivo, objetivo.replace('de', '', 1)) for n in wb.sheetnames):
            return True
        if all(f'T{i}' in wb.sheetnames for i in range(1, 9)):
            ws = wb['T1']
            for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
                if row and COLUMNAS['va'] in [str(c).strip() for c in row if c]: return True
        return False
    except Exception:
        return False

def buscar_matriz_evaluacion(raw: pathlib.Path) -> pathlib.Path:
    cands = [p for p in archivos_raw(raw) if p.suffix.lower() in ('.xlsx', '.xlsm') and es_matriz_evaluacion(p)]
    if not cands:
        sys.exit(f"No encontré la matriz de evaluación del Banco en {raw} (xlsx con la hoja de método o T1–T8 con columna VA). Pásala como segundo argumento.")
    if len(cands) > 1:
        print("⚠ Varias matrices candidatas, uso la primera:", [c.name for c in cands])
    return cands[0]
