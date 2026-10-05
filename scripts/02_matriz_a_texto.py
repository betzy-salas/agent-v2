#!/usr/bin/env python3
"""F1 · Matriz de evaluación del Banco (xlsx) → vendors/<vendor>/matriz/T1..T8.txt

Uso: python scripts/02_matriz_a_texto.py <vendor> [ruta_al_xlsx]
Si no se da la ruta, busca en vendors/<vendor>/docs/raw el xlsx que tenga la hoja de método o hojas T1–T8 con columna VA
(cabeceras configurables en scripts/_comun.py, COLUMNAS).
"""
import openpyxl, pathlib, sys
from _comun import arg_vendor, vendor_dir, buscar_matriz_evaluacion, COLUMNAS as K
v = arg_vendor(); d = vendor_dir(v)
xlsx = sys.argv[2] if len(sys.argv) > 2 else buscar_matriz_evaluacion(d / 'docs' / 'raw')
print('Matriz:', xlsx)
out = d / 'matriz'; out.mkdir(exist_ok=True)
wb = openpyxl.load_workbook(xlsx, data_only=True)
for s in ['T1','T2','T3','T4','T5','T6','T7','T8']:
    if s not in wb.sheetnames: print(s, 'NO EXISTE en el Excel'); continue
    ws = wb[s]; rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and K['numero'] in [str(c).strip() for c in r if c])
    hdr = [str(c).strip() if c else '' for c in rows[h]]
    col = lambda name: next((i for i, c in enumerate(hdr) if c.startswith(name)), None)
    iN, iCat, iQ, iCrit, iCrit2 = col(K['numero']), col(K['categoria']), col(K['pregunta']), col(K['criticidad']), col(K['criticidad_nueva'])
    iEst, iSus = col(K['estado']), col(K['sustento'])
    exact = lambda name: next((i for i, c in enumerate(hdr) if c == name), None)
    iEval, iVA, iEvdor = exact(K['eval_texto']), exact(K['va']), col(K['evaluador'])  # 'Eval' exacto: 'Evaluador…' va antes y startswith lo confundía
    iResp = [i for i, c in enumerate(hdr) if c.startswith(K['respuesta'])]
    with open(out / f'{s}.txt', 'w', encoding='utf-8') as o:
        o.write(f'# Hoja {s}\n')
        for r in rows[h+1:]:
            if not r[iN] or not str(r[iN]).startswith('E'): continue
            if iEst is not None and str(r[iEst]) in ('Eliminar', 'Omitir Respuesta'): continue
            if r[iQ] and str(r[iQ]).strip().lower().startswith(('se elimin', 'se elemin')): continue  # p.ej. T7-E32 «Se eleminó requerimiento» con estado vacío
            crit = (r[iCrit2] if iCrit2 is not None and r[iCrit2] else r[iCrit]) or ''
            resp = ' | '.join(str(r[i]) for i in iResp if r[i] and str(r[i]).strip() not in ('Sin cambios.', 'None', 'nan'))
            va = r[iVA] if iVA is not None else None
            va = int(va) if isinstance(va, (int, float)) else (va if va not in (None, '') else None)
            evaltxt = r[iEval] if iEval is not None and r[iEval] not in (None, '') else None
            evdor = r[iEvdor] if iEvdor is not None and r[iEvdor] not in (None, '') else None
            o.write(f'\n[{r[iN]}] cat={r[iCat]} crit={crit} va={va} eval={evaltxt} evaluador={evdor}\nQ: {r[iQ]}\nPROVEEDOR: {resp}\nSUSTENTO: {r[iSus] if iSus is not None else ""}\n')
    print(s, 'listo:', sum(1 for l in open(out / f'{s}.txt', encoding='utf-8') if l.startswith('[E')), 'ítems')