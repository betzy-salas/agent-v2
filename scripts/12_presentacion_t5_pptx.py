#!/usr/bin/env python3
"""Presentación T5 en PowerPoint (.pptx) con el diseño del deck corporativo (16:9, barra azul, numerales naranja).

Lee la misma base que scripts/10_presentacion_t5.py (<base>/<vendor>/9_semaforo_T5.md y 10_ficha_T5.md, <base>/COMPARATIVO_SEMAFORO_T5.md)
y escribe <base>/PRESENTACION_T5.pptx. No calcula ni redacta nada: solo maqueta lo que ya está escrito en esos archivos.
Si existen <base>/assets/logo.png y <base>/assets/portada.jpg los usa en la portada y en los separadores por vendor.

Uso:  python scripts/12_presentacion_t5_pptx.py [--base carpeta]      (por defecto vendors/)
Requiere: pip install python-pptx
"""
import datetime, importlib.util, math, re, sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent; RAIZ = AQUI.parent
spec = importlib.util.spec_from_file_location('p10', AQUI / '10_presentacion_t5.py'); p10 = importlib.util.module_from_spec(spec); spec.loader.exec_module(p10)
m = p10.m
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE
    from pptx.oxml.ns import qn
    from lxml import etree
except ImportError:
    sys.exit('Falta python-pptx: pip install python-pptx')

# ----------------------------------------------------------------------------- paleta del deck corporativo
NAVY, BLUE, NAVY_T, ORANGE, INK, GREY, LBLUE = '003B70', '00529B', '00205B', 'FF6400', '171F2A', '5F6B7A', '5B9BD5'
LEVEL_FILL = {'VERDE': 'B8D680', 'AMARILLO': 'FCC97D', 'ROJO': 'F7786E'}
PILL = {'VERDE': ('DCEBD4', '276221'), 'AMARILLO': ('FBE4CC', '9C5700'), 'ROJO': ('F5CCCC', 'B02A2A')}
TILE = {4: ('DCEBD4', '276221', 'Completamente cubiertas'), 3: ('E8F1F8', '00529B', 'Modificación menor'), 2: ('FBE4CC', '9C5700', 'Modificación mayor'), 1: ('FBE4CC', '9C5700', 'Parcialmente cubiertas'), 0: ('F5CCCC', 'B02A2A', 'No cubiertas')}
BOX_HEAD = {'g': '276221', 'r': 'B02A2A', 'o': '9C5700', 'p': '00529B', 'gr': '5F6B7A', 'n': NAVY}
FONT, MONO = 'Aptos', 'Consolas'
W, H = 13.333, 7.5
SIM = {'VERDE': '🟢', 'AMARILLO': '🟡', 'ROJO': '🔴'}

def rgb(h): return RGBColor.from_string(h)

# ----------------------------------------------------------------------------- primitivas
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=None, alpha=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:   # transparencia (python-pptx no la expone): <a:alpha val="…"/>
            clr = s.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr')); a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 100000)))
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(line_w or 0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def estimar_lineas(paras, w_in, pt):
    """Líneas aproximadas que ocupan los párrafos (lista de listas de runs) en un cuadro de w_in pulgadas a pt puntos."""
    cpl = max(10, int(w_in * 72 / (pt * 0.48)))
    return sum(max(1, math.ceil(sum(len(t) for t, *_ in p) / cpl)) for p in paras)

def texto(slide, x, y, w, h, paras, size=10, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, ajustar=True, min_size=6.5, espacio=2, bullets=False, margen=0.04):
    """paras: lista de párrafos; cada párrafo es str o lista de runs (texto, {bold, color, size, italic, font})."""
    paras = [[(p, {})] if isinstance(p, str) else p for p in paras]
    if ajustar:
        while size > min_size and estimar_lineas(paras, w - 2 * margen, size) * size * 1.22 / 72 + len(paras) * espacio / 72 > h - 2 * margen: size -= 0.5
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margen); tf.margin_top = tf.margin_bottom = Inches(margen); tf.vertical_anchor = anchor
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(espacio)
        if bullets:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(Inches(0.16)))); pPr.set('indent', str(int(-Inches(0.14))))
            bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '•')
        for t, o in runs:
            r = p.add_run(); r.text = t; f = r.font; f.name = o.get('font', font); f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.italic = o.get('italic', False)
            f.color.rgb = rgb(o.get('color', color))
    return tb, size

def tabla(slide, x, y, w, col_w, header, filas, size=8, head_size=8.2, fills=None, first_col_key=True, row_h=None, mono_cols=(), center_cols=(), bold_cols=()):
    """filas: lista de listas de celdas (str o lista de runs). fills: {(fila, col): hex}. Devuelve la forma."""
    n_r, n_c = len(filas) + 1, len(header)
    shp = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(0.3 * n_r)); t = shp.table
    tblPr = t._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None: tblPr.remove(st)
    tot = sum(col_w)
    for j, cw in enumerate(col_w): t.columns[j].width = Inches(w * cw / tot)
    def celda(c, txt, fill, color, bold, sz, fnt=FONT, align=PP_ALIGN.LEFT):
        c.fill.solid(); c.fill.fore_color.rgb = rgb(fill); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.035); c.vertical_anchor = MSO_ANCHOR.TOP
        tf = c.text_frame; tf.word_wrap = True; runs = [(txt, {})] if isinstance(txt, str) else txt
        lineas = txt.split('\n') if isinstance(txt, str) else None
        if lineas and len(lineas) > 1: runs = None
        if runs is not None:
            p = tf.paragraphs[0]; p.alignment = align
            for tx, o in runs:
                r = p.add_run(); r.text = tx; r.font.name = o.get('font', fnt); r.font.size = Pt(o.get('size', sz)); r.font.bold = o.get('bold', bold); r.font.color.rgb = rgb(o.get('color', color))
        else:
            for i, ln in enumerate(lineas):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; r = p.add_run(); r.text = ln; r.font.name = fnt; r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = rgb(color)
    for j, hd in enumerate(header): celda(t.cell(0, j), hd, NAVY, 'FFFFFF', True, head_size, align=PP_ALIGN.CENTER if j in center_cols else PP_ALIGN.LEFT)
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila):
            fill = (fills or {}).get((i - 1, j)) or ('E3E8F5' if j == 0 and first_col_key else ('F6F8FB' if i % 2 == 0 else 'FFFFFF'))
            celda(t.cell(i, j), v, fill, INK, (j == 0 and first_col_key) or j in bold_cols, size, fnt=MONO if j in mono_cols else FONT, align=PP_ALIGN.CENTER if j in center_cols else PP_ALIGN.LEFT)
    if row_h:
        for i in range(n_r): t.rows[i].height = Inches(row_h)
    return shp

def barra(slide, titulo, derecha=''):
    rect(slide, 0, 0, W, 0.64, NAVY)
    texto(slide, 0.42, 0.08, 9.6, 0.5, [[(titulo, {'bold': True, 'color': 'FFFFFF', 'size': 17})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    if derecha: texto(slide, 9.0, 0.14, 3.95, 0.4, [[(derecha, {'color': 'DCE8F6', 'size': 8})]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)

FUENTE = 'Fuente: Matriz técnica del Banco, hoja T5 · semáforo y ficha T5 de cada vendor'
def pie(slide, n, fuente=FUENTE):
    texto(slide, 0.42, 7.1, 7.5, 0.3, [[(fuente, {'color': GREY, 'size': 6.8})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 4.2, 7.18, 5, 0.3, [[('Datos elaborados por BCP para uso Interno', {'color': LBLUE, 'size': 6.8})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 12.2, 7.1, 0.75, 0.3, [[(str(n), {'color': BLUE, 'size': 7.5, 'bold': True})]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)

def seccion(slide, num, etiqueta, titulo, sub):
    c = rect(slide, 0.5, 0.36, 0.36, 0.36, ORANGE, shape=MSO_SHAPE.OVAL)
    texto(slide, 0.5, 0.36, 0.36, 0.36, [[(num, {'bold': True, 'color': 'FFFFFF', 'size': 10})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 0.98, 0.33, 6, 0.42, [[(etiqueta.upper(), {'bold': True, 'color': ORANGE, 'size': 10.5})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 0.5, 0.86, 12.3, 0.62, [[(titulo, {'bold': True, 'color': NAVY_T, 'size': 26})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    texto(slide, 0.5, 1.5, 12.3, 0.4, [[(sub, {'color': '3D4A5C', 'size': 10.5})]], ajustar=False)

def caja(slide, x, y, w, h, titulo, items, cls, size=8):
    """Caja con cabecera de color y viñetas (título en negrita, cuerpo y fuente en gris pequeño)."""
    rect(slide, x, y, w, h, None, 'E3E6EC')
    rect(slide, x, y, w, 0.3, BOX_HEAD[cls])
    texto(slide, x + 0.08, y, w - 0.16, 0.3, [[(titulo, {'bold': True, 'color': 'FFFFFF', 'size': 9})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    paras = []
    for v in items:
        tit, cuerpo, meta = partes_vineta(v); runs = []
        if tit: runs.append((tit + ' ', {'bold': True, 'color': NAVY_T}))
        runs.append((cuerpo, {}))
        if meta: runs.append(('  ' + meta, {'color': GREY, 'size': max(5.8, size - 1.4)}))
        paras.append(runs)
    if not paras: paras = [[('—', {})]]
    texto(slide, x + 0.1, y + 0.36, w - 0.2, h - 0.42, paras, size=size, bullets=True, espacio=3)

def partes_vineta(v):
    v = v.strip(); mm = re.match(r'^\*\*(.+?)\*\*\s*:?\s*(.*)$', v, re.S)
    titulo, resto = (mm.group(1), mm.group(2)) if mm else ('', v)
    partes = re.split(r'\s(?=(?:Fuente|Confianza|Impacto|Ítems?)\s*:)', resto)
    cuerpo = re.sub(r'\*\*', '', partes[0].strip()); meta = ' · '.join(p.strip().rstrip('.') for p in partes[1:])
    return titulo, cuerpo, meta

def limpio(md): return re.sub(r'\*\*(.+?)\*\*', r'\1', md or '').strip()

# ----------------------------------------------------------------------------- construcción
def construir(base, salida):
    V = p10.cargar(base)
    if not V: sys.exit(f'No encontré ningún 9_semaforo_T5.md bajo {base}.')
    por_cat, por_vendor, final = p10.leer_comentarios(base)
    cats = []
    for d in V:
        for c in d['cats']:
            if c['cat'] not in cats: cats.append(c['cat'])
    for i, d in enumerate(V):
        d['_i'] = i; d['dist_cat'], d['dist'] = p10.distribucion_items(d['ruta']); d['ficha'] = p10.leer_ficha(d['ruta'].parent / '10_ficha_T5.md')
    orden = sorted(V, key=lambda d: -(d['pct'] or 0)); pos = {d['vendor']: k + 1 for k, d in enumerate(orden)}
    sem = lambda d: {k: sum(1 for c in d['cats'] if c['color'] == k) for k in ('VERDE', 'AMARILLO', 'ROJO')}
    logo = base / 'assets' / 'logo.png'; foto = base / 'assets' / 'portada.jpg'
    MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    hoy = datetime.date.today(); mes = f'{MESES[hoy.month - 1].capitalize()} {hoy.year}'
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H); blank = prs.slide_layouts[6]
    n = [0]
    def nueva(): return prs.slides.add_slide(blank)
    def num(): n[0] += 1; return n[0]
    # 0 portada
    s = nueva(); rect(s, 0, 0, W, H, '1B2A4A')
    if foto.exists(): s.shapes.add_picture(str(foto), Inches(0), Inches(0), Inches(W), Inches(H))
    rect(s, -4.6, -2.5, 12.8, 12.8, 'E35A1E', shape=MSO_SHAPE.OVAL, alpha=0.93)
    texto(s, 0.95, 2.55, 7.5, 1.3, [[('Proyecto Transformación', {'bold': True, 'color': 'FFFFFF', 'size': 32})], [('Nuevo Core Cuentas', {'bold': True, 'color': 'FFFFFF', 'size': 32})]], ajustar=False, espacio=0)
    texto(s, 0.95, 3.85, 7.5, 0.5, [[('Revisión Técnica Resultados T5', {'bold': True, 'italic': True, 'color': 'FFFFFF', 'size': 21})]], ajustar=False)
    texto(s, 0.95, 5.35, 4, 0.4, [[(mes, {'color': 'FFFFFF', 'size': 12})]], ajustar=False)
    if logo.exists(): s.shapes.add_picture(str(logo), Inches(W - 0.38 - 1.65), Inches(H - 0.24 - 0.79), Inches(1.65))
    # 1 ranking
    s = nueva(); seccion(s, '01', 'Evaluación técnica', 'Ranking T5: capacidades de desarrollo', 'Puntaje T5 = media de las calificaciones del Banco (0–4 por ítem) ponderada por criticidad (Indispensable 3, Deseable 1), en escala 0–100 · 16 categorías · 80 ítems por vendor · nada se recalifica')
    filas, fills = [], {}
    for i, d in enumerate(orden):
        k = sem(d); c = p10.color_de(d['pct'])
        filas.append([str(pos[d['vendor']]), m.nombre(d['vendor']), str(d['pct']), d['nivel'], str(k['VERDE']), str(k['AMARILLO']), str(k['ROJO']), p10.mini_dist(d['dist'])])
        fills[(i, 2)] = LEVEL_FILL[c]; fills[(i, 3)] = LEVEL_FILL[c]
    tabla(s, 0.5, 2.05, 12.3, [0.5, 2.4, 1.2, 2.0, 0.7, 0.7, 0.7, 2.6], ['#', 'Vendor', 'Puntaje T5', 'Nivel', '🟢', '🟡', '🔴', 'Calificaciones 4 · 3 · 2 · 1 · 0'], filas, size=10, head_size=10.5, fills=fills, first_col_key=False, mono_cols=(0, 2, 7), center_cols=(0, 2, 4, 5, 6), bold_cols=(1, 3), row_h=0.42)
    texto(s, 0.5, 5.75, 12.3, 0.4, [[('Nivel por puntaje: verde ≥ 70 (Cumple; ≥ 85 Cumple plenamente) · amarillo 50–69 (Parcial) · rojo < 50 (No cumple). 🟢 🟡 🔴 = categorías del vendor en cada color.', {'color': GREY, 'size': 8})]], ajustar=False)
    pie(s, num())
    # 2 categorías × vendors
    s = nueva(); seccion(s, '02', 'Evaluación técnica', 'Puntaje por categoría', 'Cada celda es el puntaje 0–100 de la categoría para ese vendor; el color es su nivel. Vendors en orden de ranking.')
    filas, fills = [], {}
    for i, cat in enumerate(cats):
        fila = [cat]
        for j, d in enumerate(orden):
            c = next((x for x in d['cats'] if x['cat'] == cat), None); fila.append(str(c['pct']) if c else '—')
            if c: fills[(i, j + 1)] = LEVEL_FILL[c['color']]
        filas.append(fila)
    filas.append(['Puntaje T5'] + [str(d['pct']) for d in orden])
    for j, d in enumerate(orden): fills[(len(cats), j + 1)] = LEVEL_FILL[p10.color_de(d['pct'])]
    fills[(len(cats), 0)] = 'C7D1ED'
    tabla(s, 0.5, 1.95, 12.3, [3.2] + [1.5] * len(orden), ['Categoría'] + [f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])}' for d in orden], filas, size=8, head_size=8.5, fills=fills, mono_cols=tuple(range(1, len(orden) + 1)), center_cols=tuple(range(1, len(orden) + 1)), row_h=0.27)
    pie(s, num())
    # 3 cómo leer
    s = nueva(); seccion(s, '03', 'Cómo leer', 'Qué mide cada número', '')
    caja(s, 0.5, 1.7, 6.0, 3.6, 'Fuente y puntaje', ['La fuente es la calificación del Banco en la hoja T5 de la matriz técnica de cada vendor (0–4 por ítem: Completamente Cubierto 4, Modificación Menor 3, Modificación Mayor 2, Parcialmente Cubierto 1, No Cubierto 0). Nada se recalifica.', 'El puntaje de una categoría es la media de las calificaciones de sus ítems ponderada por criticidad (Indispensable 3, Deseable 1, como en el consolidado del Banco), sobre 4 y en escala 0–100. El puntaje T5 es la misma media sobre todos los ítems del vendor.'], 'n', size=10)
    caja(s, 6.8, 1.7, 6.0, 3.6, 'Colores, fichas y comentarios', ['Verde ≥ 70, amarillo 50–69, rojo < 50: son los niveles de cumplimiento oficiales del proceso.', 'Las fortalezas, debilidades, pendientes y contradicciones de cada vendor salen de su matriz (respuesta del proveedor y sustento del evaluador) y, cuando existe, de la verificación documental de T5. Cada punto nombra su ítem y su fuente; nada proviene de información externa.'], 'n', size=10)
    pie(s, num())
    # 4 vendors
    for d in orden:
        nom = m.nombre(d['vendor']); k = sem(d); c = p10.color_de(d['pct']); f = d['ficha'] or {}; r = pos[d['vendor']]
        # separador
        s = nueva(); rect(s, -2.1, -0.55, 8.6, 8.6, None, 'EEF1F4', shape=MSO_SHAPE.OVAL, line_w=30)
        if logo.exists(): s.shapes.add_picture(str(logo), Inches(0.85), Inches(3.28), Inches(3.6))
        texto(s, 7.3, 3.9, 5.5, 0.4, [[(f'Vendor #{r} de {len(V)}', {'bold': True, 'color': '0E2841', 'size': 14})]], ajustar=False)
        texto(s, 7.3, 4.25, 5.5, 0.8, [[(nom, {'bold': True, 'color': '0E2841', 'size': 34})]], ajustar=False)
        texto(s, 7.3, 5.05, 5.5, 0.4, [[(f'T5 · {d["pct"]} · {d["nivel"]}', {'color': GREY, 'size': 10})]], ajustar=False)
        # resumen ejecutivo
        s = nueva(); barra(s, f'Resumen ejecutivo {nom.upper()}', f'T5 | {d["nivel"]} | Ranking #{r} de {len(V)}')
        texto(s, 0.42, 0.78, 12.5, 0.4, [[('Lectura general', {'bold': True, 'color': BLUE, 'size': 15})]], ajustar=False)
        lectura = limpio(f.get('lectura') or d['resumen'])
        _, sz = texto(s, 0.42, 1.12, 12.5, 1.25, [lectura], size=10.2, min_size=7.5)
        for j, v in enumerate((4, 3, 2, 1, 0)):
            fill, col, nm = TILE[v]; x = 0.42 + j * 2.54
            rect(s, x, 2.5, 2.34, 0.72, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            texto(s, x + 0.12, 2.5, 0.9, 0.72, [[(str(d['dist'].get(v, 0)), {'bold': True, 'color': col, 'size': 22})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
            texto(s, x + 0.95, 2.5, 1.35, 0.72, [[(nm, {'bold': True, 'color': col, 'size': 8.6})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
        def titulos(items, n_max=8):
            out = []
            for v in items[:n_max]:
                t, cuerpo, _ = partes_vineta(v); out.append(t or cuerpo[:160])
            return '\n'.join('• ' + t for t in out) or '—'
        pend = '\n'.join('• ' + partes_vineta(v)[1] for v in (f.get('pendientes') or [])[:4]) or '—'
        contr = f.get('contradicciones') or []
        codigos = sorted({cc for v in contr for cc in p10.codigos(v)}, key=lambda cc: int(cc[1:]))
        contr_txt = (f'{len(contr)} entre la calificación del Banco y la verificación documental: ' + ', '.join(codigos) + '. Detalle en la hoja de pendientes y contradicciones.') if contr else 'Sin contradicciones registradas en la ficha.'
        res = [(f'{d["pct"]} · {d["nivel"]}', {'bold': True, 'size': 10}), ('\n', {}), (f'🟢 {k["VERDE"]} · 🟡 {k["AMARILLO"]} · 🔴 {k["ROJO"]} de 16 categorías', {'size': 8}), ('\n', {}), ('Contradicciones: ' + contr_txt, {'size': 7.4})]
        fo, de, pe = f.get('fortalezas') or [], f.get('debilidades') or [], f.get('pendientes') or []
        tabla(s, 0.42, 3.42, 12.5, [1.8, 3.4, 3.4, 3.9], ['Resultado T5', f'Fortalezas dominantes ({len(fo)})', f'Debilidades dominantes ({len(de)})', f'Pendientes clave ({len(pe)})'],
              [[res, titulos(fo), titulos(de), pend]], size=7.8, fills={(0, 0): LEVEL_FILL[c]}, first_col_key=False)
        pie(s, num())
        # matriz por categoría
        s = nueva(); barra(s, f'Matriz por categoría · {nom}', 'Puntaje, calificaciones del Banco y lectura por categoría')
        filas, fills = [], {}
        for i, x in enumerate(d['cats']):
            dist = d['dist_cat'].get(x['cat'], {}); dem, falta = (f.get('tabla') or {}).get(x['cat'], ('', ''))
            filas.append([x['cat'], str(x['pct']), p10.mini_dist(dist), dem or d['conclusiones'].get(x['cat'], ''), falta, x['color'].capitalize()]); fills[(i, 5)] = LEVEL_FILL[x['color']]
        tabla(s, 0.42, 0.82, 12.5, [2.2, 0.75, 1.15, 3.8, 3.8, 0.9], ['Categoría', 'Puntaje', '4 · 3 · 2 · 1 · 0', 'Lo que demuestra', 'Lo que falta', 'Resultado'], filas, size=7, head_size=7.6, fills=fills, mono_cols=(1, 2), center_cols=(1,), bold_cols=(5,), row_h=0.34)
        pie(s, num())
        if f:
            s = nueva(); barra(s, f'Fortalezas y debilidades · {nom}', f'T5 | {d["pct"]} · {d["nivel"]}')
            caja(s, 0.42, 0.85, 6.15, 6.1, f'Fortalezas ({len(fo)})', fo, 'g', size=8); caja(s, 6.77, 0.85, 6.15, 6.1, f'Debilidades ({len(de)})', de, 'r', size=8)
            pie(s, num())
            s = nueva(); barra(s, f'Pendientes, contradicciones y declarado sin sustento · {nom}', f'T5 | {d["pct"]} · {d["nivel"]}')
            ss = f.get('sin_sustento') or []; h_alto = 4.3 if ss else 6.1
            caja(s, 0.42, 0.85, 6.15, h_alto, f'Pendientes ({len(pe)})', pe, 'o', size=8); caja(s, 6.77, 0.85, 6.15, h_alto, f'Contradicciones ({len(f.get("contradicciones") or [])})', f.get('contradicciones') or [], 'p', size=8)
            if ss: caja(s, 0.42, 5.3, 12.5, 1.65, f'Declarado sin sustento ({len(ss)})', ss, 'gr', size=7.8)
            pie(s, num())
    # 5 comentarios por categoría
    if any(por_cat.get(c) for c in cats):
        grupos = [cats[i:i + 6] for i in range(0, len(cats), 6)]
        for g, grupo in enumerate(grupos, 1):
            s = nueva(); barra(s, f'Comentarios por categoría ({g} de {len(grupos)})', 'Dónde coinciden y dónde se separan los vendors')
            texto(s, 0.42, 0.74, 12.5, 0.3, [[('Los puntos muestran el color de cada vendor en la categoría, en orden de ranking: ' + ' · '.join(f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])}' for d in orden) + '.', {'color': GREY, 'size': 8.6})]], ajustar=False)
            for j, cat in enumerate(grupo):
                col, fila = j % 2, j // 2; x = 0.42 + col * 6.4; y = 1.12 + fila * 1.98
                dots = ' '.join(SIM[c['color']] for d in orden for c in d['cats'] if c['cat'] == cat)
                texto(s, x, y, 4.6, 0.3, [[(cat, {'bold': True, 'color': NAVY_T, 'size': 9.6})]], ajustar=False)
                texto(s, x + 4.6, y, 1.5, 0.3, [[(dots, {'size': 8})]], align=PP_ALIGN.RIGHT, ajustar=False)
                texto(s, x, y + 0.3, 6.1, 1.6, [limpio(por_cat.get(cat, ''))], size=8.8, min_size=7)
                ln = s.shapes.add_connector(1, Inches(x), Inches(y + 1.9), Inches(x + 6.1), Inches(y + 1.9)); ln.line.color.rgb = rgb('E3E6EC'); ln.line.width = Pt(0.75)
            pie(s, num())
    # 6 resultado final
    s = nueva(); barra(s, 'Resultado final', 'Ranking por puntaje T5')
    rect(s, 0.42, 0.85, 6.15, 2.75, None, 'E3E6EC')
    x0, ancho = 2.3, 3.9
    for v in (0, 25, 50, 70, 85, 100):
        px = x0 + ancho * v / 100
        ln = s.shapes.add_connector(1, Inches(px), Inches(1.0), Inches(px), Inches(3.3)); ln.line.color.rgb = rgb('E3E2DE'); ln.line.width = Pt(0.5)
        texto(s, px - 0.25, 3.3, 0.5, 0.22, [[(str(v), {'color': GREY, 'size': 7})]], align=PP_ALIGN.CENTER, ajustar=False)
    for i, d in enumerate(orden):
        y = 1.02 + i * 0.37
        texto(s, 0.45, y, 1.8, 0.3, [[(m.nombre(d['vendor']), {'bold': True, 'size': 9})]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
        rect(s, x0, y + 0.04, ancho * d['pct'] / 100, 0.22, LEVEL_FILL[p10.color_de(d['pct'])], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        texto(s, x0 + ancho * d['pct'] / 100 + 0.05, y, 2.2, 0.3, [[(f'{d["pct"]} · {d["nivel"]}', {'size': 8, 'font': MONO})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    filas, fills = [], {}
    for i, d in enumerate(orden):
        k = sem(d); c = p10.color_de(d['pct']); filas.append([str(pos[d['vendor']]), m.nombre(d['vendor']), str(d['pct']), d['nivel'], str(k['VERDE']), str(k['AMARILLO']), str(k['ROJO'])]); fills[(i, 2)] = LEVEL_FILL[c]; fills[(i, 3)] = LEVEL_FILL[c]
    tabla(s, 0.42, 3.8, 6.15, [0.4, 1.9, 0.9, 1.7, 0.5, 0.5, 0.5], ['#', 'Vendor', 'Puntaje', 'Nivel', '🟢', '🟡', '🔴'], filas, size=8, fills=fills, first_col_key=False, mono_cols=(0, 2), center_cols=(0, 2, 4, 5, 6), bold_cols=(1,), row_h=0.32)
    caja(s, 6.77, 0.85, 6.15, 5.0, 'Comentario final', [limpio(final) or 'Pendiente de redacción.'], 'p', size=9.2)
    pie(s, num())
    # 7 lo pendiente
    s = nueva(); barra(s, 'Lo pendiente', 'Categorías en rojo y amarillo por vendor')
    texto(s, 0.42, 0.74, 12.5, 0.3, [[('Lo que cada vendor dejó sin cubrir o sin sustento según la calificación del Banco; el detalle por ítem está en el semáforo y la ficha de cada vendor.', {'color': GREY, 'size': 9})]], ajustar=False)
    for i, d in enumerate(orden):
        col, fila = i % 3, i // 3; x = 0.42 + col * 4.2; y = 1.15 + fila * 2.95
        items = [f'{SIM[c["color"]]} {c["cat"]} · {c["pct"]}' for c in sorted([c for c in d['cats'] if c['color'] in ('ROJO', 'AMARILLO')], key=lambda c: c['pct'] or 0)] or ['Ninguna categoría en rojo ni amarillo.']
        caja(s, x, y, 4.0, 2.8, f'#{pos[d["vendor"]]} {m.nombre(d["vendor"])} · categorías en rojo y amarillo', items, 'n', size=8.2)
    pie(s, num(), f'Presentación generada el {hoy.strftime("%d-%m-%Y")} · ' + FUENTE)
    # 8 estrategia de implementación (general + por vendor), si existen los informes de implementación
    IMP = p10.datos_implementacion(V); g = IMP['general']
    if g:
        FI = 'Fuente: rúbrica de implementación (verif/IMPL.md), informe de implementación (7_) y COMPARATIVO_IMPLEMENTACION.md'
        hdr, rows = g['resumen']
        if hdr:
            col = {m._norm(h): i for i, h in enumerate(hdr)}
            def ci(*pref): return next((i for h, i in col.items() if any(h.startswith(m._norm(p)) for p in pref)), None)
            idx = [ci('vendor'), ci('enfoque'), ci('duraci'), ci('implementador'), ci('% estrategia', 'ei'), ci('% fabricante'), ci('% implementador'), ci('señales', 'senales')]
            filas, fills = [], {}
            for i, r in enumerate(rows):
                r = r + [''] * (len(hdr) - len(r)); fila = [(m.nombre(r[j]) if k == 0 else r[j]) if j is not None else '' for k, j in enumerate(idx)]
                for k in (4, 5, 6):
                    mm = re.search(r'[\d.]+', fila[k] or ''); fills[(i, k)] = LEVEL_FILL.get(p10.nivel_pct(float(mm.group(0)) if mm else None), 'EEF1F4')
                filas.append(fila)
            s = nueva(); seccion(s, '04', 'Estrategia de implementación', 'Cómo propone implantar cada vendor', m.nombres_en_texto(g['titular']))
            tabla(s, 0.5, 2.1, 12.3, [1.4, 2.6, 1.5, 2.4, 1.0, 1.0, 1.1, 0.9], ['Vendor', 'Enfoque de despliegue', 'Duración propuesta', 'Implementador', '% Estrategia (EI)', '% Fabricante (F)', '% Implementador (I)', 'Señales de alerta'], filas, size=8.2, head_size=8, fills=fills, mono_cols=(4, 5, 6, 7), center_cols=(4, 5, 6, 7), bold_cols=(4, 5, 6), row_h=0.42)
            texto(s, 0.5, 5.6, 12.3, 0.5, [[('Rúbrica corporativa de implementación (0–5 por criterio, ponderada): EI estrategia de implementación · F perfil del fabricante · I perfil del implementador. Color por nivel: verde ≥ 70 · amarillo 50–69 · rojo < 50. Fuente: verificación de implementación y comparativo de implementación de cada vendor.', {'color': GREY, 'size': 8})]], ajustar=False)
            pie(s, num(), FI)
        hdr, rows = g['ei']
        if hdr:
            filas, fills = [], {}
            for i, r in enumerate(rows):
                r = r + [''] * (len(hdr) - len(r)); fila = [r[0]]
                for j, c in enumerate(r[1:], 1):
                    partes = c.split(); fila.append([(partes[0] + ' ', {'bold': True, 'font': MONO}), (' '.join(partes[1:]), {})] if partes else '—')
                    if partes: fills[(i, j)] = LEVEL_FILL.get(p10.color_0_5(partes[0]), 'EEF1F4')
                filas.append(fila)
            s = nueva(); barra(s, 'Estrategia de implementación · criterio a criterio', 'Rúbrica EI-01 a EI-12, puntaje 0–5 con su descriptor')
            tabla(s, 0.42, 0.82, 12.5, [2.4] + [1.68] * (len(hdr) - 1), ['Criterio EI'] + [m.nombre(h) for h in hdr[1:]], filas, size=7.4, head_size=7.8, fills=fills, row_h=0.3)
            if g['diferencias']: caja(s, 0.42, 4.95, 12.5, 2.0, 'Criterios con mayor diferencia y por qué', [m.nombres_en_texto(x) for x in g['diferencias']], 'n', size=8)
            pie(s, num(), 'Fuente: rúbrica de implementación de cada vendor (verif/IMPL.md) y COMPARATIVO_IMPLEMENTACION.md')
        if g['transversal']:
            s = nueva(); barra(s, 'Lo que cambia entre vendors en la implementación', 'Comparación transversal con la fuente documental de cada vendor')
            for j, (t, fu) in enumerate(g['transversal']):
                colm, fila = j % 2, j // 2; x = 0.42 + colm * 6.4; y = 0.85 + fila * 2.05
                texto(s, x, y, 6.1, 1.95, [[(str(j + 1), {'bold': True, 'color': NAVY_T, 'size': 10})], [(t, {'size': 8.8})], [('Fuentes: ' + fu + '.', {'color': GREY, 'size': 6.6})]], size=8.8, min_size=6.2, espacio=3)
                ln = s.shapes.add_connector(1, Inches(x), Inches(y + 1.98), Inches(x + 6.1), Inches(y + 1.98)); ln.line.color.rgb = rgb('E3E6EC'); ln.line.width = Pt(0.75)
            pie(s, num(), 'Fuente: informes de implementación (7_), rúbricas (IMPL) y análisis de despliegue (8_) de cada vendor')
        DESP, PLAT = p10.datos_despliegue(orden)
        if DESP:
            s = nueva(); barra(s, 'Modelos de despliegue · plataformas por modelo', 'Resumen por vendor: una fila por modelo y una columna por nube o plataforma')
            texto(s, 0.42, 0.74, 12.5, 0.3, [[('Solo marca lo que los documentos del vendor asocian a cada modelo de despliegue; la fuente de cada fila está en el análisis de despliegue del vendor.', {'size': 9})]], ajustar=False)
            def punto(slide, cx, cy, est, r=0.065):
                if est == 'vacio': return rect(slide, cx - 0.035, cy - 0.035, 0.07, 0.07, 'E3E6EC', shape=MSO_SHAPE.OVAL)
                o = rect(slide, cx - r, cy - r, 2 * r, 2 * r, '2A78D6' if est == 'of' else None, '2A78D6' if est in ('of', 'cond') else '8A8F98', shape=MSO_SHAPE.OVAL, line_w=1.1)
                if est == 'no': o.line.dash_style = MSO_LINE.DASH
                return o
            partes = [('Modelo: ', {'bold': True, 'color': GREY, 'size': 7})]
            for k, c, t in p10.LEYENDA_CHIPS: partes += [(c, {'bold': True, 'color': {'si': BLUE, 'cond': '9C5700', 'tiene': '8A4B00'}.get(k, GREY), 'size': 7}), (f' = {t} · ', {'color': GREY, 'size': 7})]
            texto(s, 0.42, 0.9, 12.5, 0.22, [partes], anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
            xl = 0.42
            for k, t in zip(('of', 'cond', 'no', 'vacio'), p10.LEYENDA_NUBES):
                punto(s, xl + 0.07, 1.26, k); texto(s, xl + 0.17, 1.12, 4.0, 0.28, [[(t, {'color': GREY, 'size': 7})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0); xl += 0.26 + len(t) * 0.048
            texto(s, xl + 0.1, 1.12, 12.9 - xl, 0.28, [[(p10.NOTA_NUBES, {'color': GREY, 'size': 6.4})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
            CHIP = {'si': ('E8F1F8', BLUE), 'no': ('EEF1F4', GREY), 'cond': ('FBE4CC', '9C5700'), 'tiene': ('FFFFFF', '8A4B00'), '': ('EEF1F4', GREY)}; notas = m.notas_despliegue()
            for q, (vend, i, filas) in enumerate(DESP[:6]):
                colm, fila = q % 3, q // 3; x = 0.42 + colm * 4.2; y = 1.42 + fila * 2.8; w, h = 4.0, 2.72; nota = notas.get(m._norm(vend), '')
                rect(s, x, y, w, h, None, 'E3E6EC'); rect(s, x, y, w, 0.06, m.vcolor(i).lstrip('#'))
                texto(s, x + 0.1, y + 0.1, w - 0.2, 0.28, [[(m.nombre(vend).upper(), {'bold': True, 'color': NAVY_T, 'size': 9.5})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
                wl = 1.3; wc = (w - 0.2 - wl) / len(PLAT); yh = y + 0.42
                rect(s, x + 0.1, yh, w - 0.2, 0.22, 'F6F8FB')
                texto(s, x + 0.12, yh, wl, 0.22, [[('MODELO', {'bold': True, 'color': GREY, 'size': 5.6})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
                for j, pn in enumerate(PLAT): texto(s, x + 0.1 + wl + j * wc - 0.05, yh, wc + 0.1, 0.22, [[(pn.upper(), {'bold': True, 'color': GREY, 'size': 5.4})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
                rh = (h - 0.72 - (0.52 if nota else 0)) / len(filas)
                for r_i, (nm, cls, corto, marcas) in enumerate(filas):
                    ry = yh + 0.24 + r_i * rh
                    texto(s, x + 0.12, ry, wl, 0.17, [[(nm, {'bold': True, 'size': 7})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
                    fill, colr = CHIP.get(cls, CHIP['']); wchip = min(wl - 0.05, 0.16 + len(corto) * 0.045)
                    ch = rect(s, x + 0.12, ry + 0.18, wchip, 0.15, fill, 'C97A1A' if cls == 'tiene' else None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=0.75)
                    if cls == 'tiene': ch.line.dash_style = MSO_LINE.DASH
                    texto(s, x + 0.12, ry + 0.18, wchip, 0.15, [[(corto, {'bold': True, 'color': colr, 'size': 5.6})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ajustar=False, margen=0)
                    for j, pn in enumerate(PLAT): punto(s, x + 0.1 + wl + j * wc + wc / 2, ry + rh / 2 - 0.01, marcas.get(pn, 'vacio'))
                    if r_i < len(filas) - 1:
                        ln = s.shapes.add_connector(1, Inches(x + 0.1), Inches(ry + rh - 0.01), Inches(x + w - 0.1), Inches(ry + rh - 0.01)); ln.line.color.rgb = rgb('EEF1F4'); ln.line.width = Pt(0.5)
                if nota:   # nota del vendor (vendors/NOTAS_DESPLIEGUE.md): nubes en firme, compatibilidad declarada y contradicciones
                    yn = yh + 0.24 + len(filas) * rh + 0.02
                    ln = s.shapes.add_connector(1, Inches(x + 0.1), Inches(yn), Inches(x + w - 0.1), Inches(yn)); ln.line.color.rgb = rgb('E3E6EC'); ln.line.width = Pt(0.5); ln.line.dash_style = MSO_LINE.DASH
                    texto(s, x + 0.08, yn + 0.02, w - 0.16, y + h - yn - 0.06, [[(m.nombres_en_texto(nota), {'color': '3D4A5C', 'size': 6})]], size=6, min_size=5, espacio=0, margen=0.03)
            pie(s, num(), 'Fuente: análisis de despliegue de cada vendor (8_) · informe de comparativos, hoja 4')
        for d in orden:
            x = IMP['vendors'].get(d['vendor'])
            if not x: continue
            nom = m.nombre(d['vendor']); rs = x.get('resumen') or {}; pv = g['por_vendor'].get(d['vendor'], {})
            s = nueva(); barra(s, f'Estrategia de implementación · {nom}', 'Ficha resumen del informe de implementación')
            y0 = 0.78
            if x.get('titular'):
                texto(s, 0.42, y0, 12.5, 0.6, [[('En una frase. ', {'bold': True, 'color': BLUE}), (m.nombres_en_texto(x['titular']), {})]], size=9.6, min_size=7.5); y0 += 0.66
            for j, (k, nm) in enumerate((('EI', '% Estrategia de implementación'), ('F', '% Perfil del fabricante'), ('I', '% Perfil del implementador'))):
                p = rs.get(k); lvl = p10.nivel_pct(p); fill, colr = PILL.get(lvl, ('EEF1F4', GREY)); xx = 0.42 + j * 2.54
                rect(s, xx, y0, 2.34, 0.72, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                texto(s, xx + 0.1, y0, 1.0, 0.72, [[(str(p) if p is not None else '—', {'bold': True, 'color': colr, 'size': 20})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
                texto(s, xx + 1.05, y0, 1.25, 0.72, [[(nm, {'bold': True, 'color': colr, 'size': 8})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
            rect(s, 0.42 + 3 * 2.54, y0, 12.5 - 3 * 2.54, 0.72, 'EEF1F4', shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            texto(s, 0.42 + 3 * 2.54 + 0.1, y0, 12.3 - 3 * 2.54, 0.72, [[('Implementador(es)', {'bold': True, 'color': NAVY_T, 'size': 9})], [('; '.join(x.get('implementadores') or ['—']), {'size': 8.4})]], anchor=MSO_ANCHOR.MIDDLE, size=8.4, min_size=6.5)
            y0 += 0.86
            filas, fills = [], {}
            for i, (e, qf, fu, c) in enumerate(p10.ficha_impl_filas(x)):
                filas.append([e, qf, [(fu, {'color': GREY, 'size': 6.6})], c]); lvl = {'ALTA': 'VERDE', 'MEDIA': 'AMARILLO', 'BAJA': 'ROJO'}.get(c.strip().upper())
                if lvl: fills[(i, 3)] = LEVEL_FILL[lvl]
            tabla(s, 0.42, y0, 12.5, [1.9, 5.4, 4.3, 0.9], ['Elemento', 'Qué propone', 'Fuente', 'Confianza'], filas, size=7.6, head_size=7.8, fills=fills, center_cols=(3,), bold_cols=(3,), row_h=0.44)
            pie(s, num(), f'Fuente: informe de implementación (7_) y rúbrica de implementación (verif/IMPL.md) de {nom}')
            s = nueva(); barra(s, f'Pros, contras y veredicto de implementación · {nom}', 'Del informe de implementación y del comparativo entre vendors')
            def it_box(items, lead):
                return [f'**{i["afirmacion"]}**: {lead} {i.get("importa", "")} Fuente: {i["fuente"]} · Confianza: {i["confianza"]}' if i.get('importa') else f'**{i["afirmacion"]}**. Fuente: {i["fuente"]} · Confianza: {i["confianza"]}' for i in items]
            abajo = [(t, next((v for k, v in pv.items() if m._norm(cl) in m._norm(k)), []), cls) for t, cl, cls in (('Veredicto en una frase', 'Veredicto', 'p'), ('Señales de alerta', 'Señales', 'o'), ('Lo que deja al Banco', 'deja al Banco', 'gr'))]
            abajo = [a for a in abajo if a[1]]; h_sup = 4.3 if abajo else 6.1
            caja(s, 0.42, 0.85, 6.15, h_sup, 'Pros (5)', it_box((x.get('pros_items') or [])[:5], 'Por qué importa:'), 'g', size=7.8)
            caja(s, 6.77, 0.85, 6.15, h_sup, 'Contras (5)', it_box((x.get('contras_items') or [])[:5], 'Consecuencia:'), 'r', size=7.8)
            if abajo:
                wcaja = (12.5 - 0.2 * (len(abajo) - 1)) / len(abajo)
                for j, (t, lst, cls) in enumerate(abajo): caja(s, 0.42 + j * (wcaja + 0.2), 5.3, wcaja, 1.65, t, [m.nombres_en_texto(z) for z in lst], cls, size=7.8)
            pie(s, num(), f'Fuente: informe de implementación (7_) de {nom} y COMPARATIVO_IMPLEMENTACION.md')
    prs.save(salida); return len(prs.slides)

def main():
    args = sys.argv[1:]
    base = Path(args[args.index('--base') + 1]) if '--base' in args else RAIZ / 'vendors'
    if not base.is_absolute(): base = RAIZ / base
    salida = base / 'PRESENTACION_T5.pptx'
    print(f'PPTX: {salida} · {construir(base, salida)} diapositivas')

if __name__ == '__main__':
    main()
