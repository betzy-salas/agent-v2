#!/usr/bin/env python3
"""Lámina independiente «Modelos de despliegue y nubes»: resumen por vendor (qué modelo ofrece: SaaS, PaaS, nube privada,
on-premise, híbrido; nubes ofrecidas al Banco; mencionadas pero no ofrecidas; quién opera) y cuántos vendors ofrecen cada nube
en firme. Es el resumen de las hojas 2 y 5 del informe de comparativos (scripts/08_comparativos_pdf.py), en una sola lámina
para adjuntar a un correo. Mismo diseño que las láminas del script 13 (deck corporativo 16:9).

Fuentes (nada se redacta ni se calcula de nuevo): vendors/COMPARATIVO_DESPLIEGUE.md (tabla comparativa por proveedor, del
agente analista-despliegue) y el análisis de despliegue 8_modelo_despliegue.md de cada vendor (conteo de nubes, igual que la
hoja 5 del informe de comparativos). Los vendors van en el orden del ranking T5.

Uso:  python scripts/15_lamina_despliegue.py [--base carpeta]
   → <base>/LAMINA_3_despliegue.html, .pdf, LAMINA_3_despliegue.png (2560×1440) y LAMINA_3_despliegue.pptx
Requiere: Chrome (PDF), PyMuPDF (PNG) y python-pptx (PPTX); si falta alguno, genera lo que pueda y lo avisa.
"""
import datetime, html, importlib.util, sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent; RAIZ = AQUI.parent
def _cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, AQUI / archivo); mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod
p10 = _cargar('p10', '10_presentacion_t5.py'); p08 = _cargar('p08', '08_comparativos_pdf.py'); p13 = _cargar('p13', '13_laminas_correo.py'); m = p10.m
esc = html.escape

CSS = p13.CSS + """
.l3 table.t{font-size:8.4pt} .l3 table.t td{padding:5pt 6pt;line-height:1.32;vertical-align:top} .l3 table.t th{padding:5pt 6pt;font-size:8.2pt}
.l3 .chipm{margin-left:0;font-size:7.8pt;padding:1.5pt 7pt} .l3 .resto{display:block;color:#5f6b7a;font-size:7.2pt;line-height:1.3;margin-top:2pt}
.cnt{display:flex;gap:12pt;margin-top:8pt} .cnt .ct{flex:1;background:#eef1f4;border-radius:8pt;padding:11pt 14pt;min-width:0}
.cnt .ct b{display:block;font-family:"JetBrains Mono",Consolas,monospace;font-size:24pt;font-weight:800;color:#00205b;line-height:1.1} .cnt .ct b small{font-family:inherit;font-size:10pt;font-weight:700;color:#5f6b7a}
.cnt .ct span{display:block;font-size:9.4pt;font-weight:700;color:#00529b;text-transform:uppercase;letter-spacing:.04em;margin-top:2pt}
.cnt .ct.cero b{color:#8a8f98}
"""
COLS = (('saas', 'SaaS'), ('paas', 'PaaS'), ('nube privada', 'Nube privada'), ('on-premise', 'On-premise'), ('hibrido', 'Híbrido'),
        ('nubes ofrecidas', 'Nubes ofrecidas para el Banco'), ('mencionadas', 'Nubes mencionadas, no ofrecidas o excluidas'), ('quien opera', 'Quién opera'))

def datos(base):
    """Vendors en orden de ranking T5 con su fila de la tabla comparativa por proveedor y el conteo de vendors por nube."""
    V = p10.cargar(base)
    if not V: sys.exit(f'No encontré ningún 9_semaforo_T5.md bajo {base}.')
    for i, d in enumerate(V): d['_i'] = i
    orden = sorted(V, key=lambda d: -(d['pct'] or 0)); pos = {d['vendor']: k + 1 for k, d in enumerate(orden)}
    comp = m.VENDORS / 'COMPARATIVO_DESPLIEGUE.md'
    if not comp.exists(): sys.exit(f'Falta {comp} (lo escribe el agente analista-despliegue).')
    hdr, rows = m.first_table(m.find_section(m.sections(m.leer_md(comp)), 'por proveedor'))
    idx = {}
    for clave, _ in COLS:
        idx[clave] = next((i for i, h in enumerate(hdr) if m._norm(h).replace('í', 'i').startswith(clave) or clave in m._norm(h)), None)
    filas = []
    for d in orden:
        r = next((r for r in rows if m._norm(r[0]) == m._norm(d['vendor']) or m._norm(m.nombre(r[0])) == m._norm(m.nombre(d['vendor']))), None)
        celdas = {clave: (m.nombres_en_texto(r[idx[clave]]).strip() if r and idx[clave] is not None and idx[clave] < len(r) else '—') for clave, _ in COLS}
        filas.append((d['vendor'], pos[d['vendor']], celdas))
    DESP, _ = p10.datos_despliegue(orden)
    conteo = [(pn, sum(1 for _, _, fl in DESP if any(marcas.get(pn) == 'of' for _, _, _, marcas in fl))) for pn, _ in p08.PLATAFORMAS]
    return filas, conteo, len(DESP)

def chip_partes(txt):
    cls, corto = m.chip_modelo(txt); t = (txt or '').strip()
    resto = t[len(corto):].strip(' ():;') if t.lower().startswith(corto.lower()) else t
    corto = {'si': 'Sí', 'no': 'No', 'cond': 'Condicionado'}.get(cls, corto if corto != '—' else 'N/E')
    if m._norm(t) in ('n/e', 'ne', '—', '-', ''): corto, cls, resto = 'N/E', '', ''
    return cls, corto, resto

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
TITULO = 'Revisión Técnica Resultados T5 · Modelos de despliegue y nubes'
LEYENDA = ('Sí = el vendor ofrece el modelo; Condicionado = lo ofrece con condiciones o solo declara compatibilidad; No = lo excluye; N/E = no lo especifica en sus documentos. '
           'Cada celda tiene su cita documental en el análisis de despliegue del vendor; «en firme» = ofrecida al Banco para al menos un modelo.')
FUENTE = 'Fuente: comparativo y análisis de despliegue (8_) de cada vendor, sobre su propuesta y anexos'

def construir_html(base, filas, conteo, n):
    logo = p10._data_uri(base / 'assets' / 'logo.png'); t = datetime.date.today(); mes = f'{MESES[t.month - 1].capitalize()} {t.year}'
    def bar(titulo, derecha=''): return f'<div class="bar"><h2>{esc(titulo)}</h2><span class="r">{esc(derecha)}</span></div>'
    def pie(fuente): return f'<div class="pie"><span>{esc(fuente)}</span><span class="c">Datos elaborados por BCP para uso Interno</span>' + (f'<img src="{logo}" alt="">' if logo else '') + '</div>'
    def chip(txt):
        cls, corto, resto = chip_partes(txt)
        return f'<span class="chipm {cls or "x"}">{esc(corto)}</span>' + (f'<span class="resto">{esc(resto)}</span>' if resto else '')
    trs = ''
    for vend, p, c in filas:
        trs += f'<tr><td class="k">#{p} {esc(m.nombre(vend))}</td>' + ''.join(f'<td>{chip(c[k])}</td>' for k, _ in COLS[:5]) + ''.join(f'<td>{esc(c[k])}</td>' for k, _ in COLS[5:]) + '</tr>'
    anchos = ['9%', '8%', '8%', '9%', '8%', '8%', '15%', '20%', '']
    th = ''.join(f'<th style="width:{a}">{esc(t)}</th>' for a, t in zip(anchos, ['Vendor (ranking T5)'] + [t for _, t in COLS]))
    tiles = ''.join(f'<div class="ct{" cero" if k == 0 else ""}"><b>{k} <small>de {n}</small></b><span>{esc(pn)}</span></div>' for pn, k in conteo)
    H = [f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Lámina despliegue y nubes</title><style>{CSS}</style></head><body>',
         f'<section class="s l3">{bar(TITULO, "Qué modelo ofrece cada vendor y en qué nubes · " + mes)}<div class="body">'
         f'<table class="t"><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>'
         f'<p class="leg" style="margin:6pt 0 8pt;font-size:8pt">{esc(LEYENDA)}</p>'
         f'<h3 class="az" style="font-size:15pt;margin-top:4pt">Cuántos vendors ofrecen cada nube en firme</h3><div class="cnt">{tiles}</div></div>' + pie(FUENTE) + '</section>',
         '</body></html>']
    return '\n'.join(H)

def construir_pptx(base, filas, conteo, n, salida):
    try:
        p12 = _cargar('p12', '12_presentacion_t5_pptx.py')
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    except (ImportError, SystemExit):
        print('AVISO: sin python-pptx no se genera el .pptx (pip install python-pptx)'); return 0
    p14 = _cargar('p14', '14_laminas_correo_pptx.py')
    W, H = p12.W, p12.H; NAVY_T, BLUE, GREY, INK = p12.NAVY_T, p12.BLUE, p12.GREY, p12.INK
    t = datetime.date.today(); mes = f'{MESES[t.month - 1].capitalize()} {t.year}'
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H); s = prs.slides.add_slide(prs.slide_layouts[6])
    p12.barra(s, TITULO, 'Qué modelo ofrece cada vendor y en qué nubes · ' + mes)
    COLOR = {'si': BLUE, 'no': GREY, 'cond': '9C5700', '': GREY}
    def celda_chip(txt):
        cls, corto, resto = chip_partes(txt)
        return [(corto, {'bold': True, 'color': COLOR[cls]})] + ([(' ' + resto, {'color': GREY, 'size': 6.4})] if resto else [])
    cuerpo = [[f'#{p} {m.nombre(vend)}'] + [celda_chip(c[k]) for k, _ in COLS[:5]] + [c[k] for k, _ in COLS[5:]] for vend, p, c in filas]
    col_w = [1.15, 1.0, 1.0, 1.15, 1.0, 1.0, 1.9, 2.5, 1.8]; sz = 7.8; y0 = 0.8
    shp = p12.tabla(s, 0.42, y0, 12.5, col_w, ['Vendor (ranking T5)'] + [t for _, t in COLS], cuerpo, size=sz, head_size=8.4)
    tot = sum(col_w); ws = [12.5 * c / tot for c in col_w]
    y = y0 + 0.36; shp.table.rows[0].height = Inches(0.36)
    for i, fila in enumerate(cuerpo, 1):
        lineas = max(p12.estimar_lineas([[(c, {})] if isinstance(c, str) else c], ws[j] - 0.12, sz) for j, c in enumerate(fila))
        h = lineas * sz * 1.3 / 72 + 0.14; shp.table.rows[i].height = Inches(h); y += h
    p12.texto(s, 0.42, y + 0.04, 12.5, 0.34, [[(LEYENDA, {'color': GREY, 'size': 8})]], ajustar=False)
    y_h = y + 0.5; p14.h3(s, 0.42, y_h, 12.5, 'Cuántos vendors ofrecen cada nube en firme', size=14)
    y_t = y_h + 0.5; alto = min(1.0, 7.0 - y_t); gap = 0.14; w = (12.5 - gap * (len(conteo) - 1)) / len(conteo)
    for j, (pn, k) in enumerate(conteo):
        x = 0.42 + j * (w + gap); p12.rect(s, x, y_t, w, alto, 'EEF1F4')
        p12.texto(s, x + 0.12, y_t + 0.1, w - 0.2, 0.45, [[(str(k), {'bold': True, 'color': NAVY_T if k else '8A8F98', 'size': 22, 'font': p12.MONO}), (f' de {n}', {'bold': True, 'color': GREY, 'size': 10, 'font': p12.MONO})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
        p12.texto(s, x + 0.12, y_t + 0.58, w - 0.2, 0.3, [[(pn.upper(), {'bold': True, 'color': BLUE, 'size': 9.2})]], anchor=MSO_ANCHOR.MIDDLE, ajustar=False)
    p14.pie(s, FUENTE, base / 'assets' / 'logo.png')
    if y_t + alto > 7.05: print(f'AVISO: el contenido llega a {y_t + alto:.2f} in y pisa el pie (7.1 in).')
    prs.save(salida); return 1

def main():
    args = sys.argv[1:]
    base = Path(args[args.index('--base') + 1]) if '--base' in args else RAIZ / 'vendors'
    if not base.is_absolute(): base = RAIZ / base
    filas, conteo, n = datos(base)
    out = base / 'LAMINA_3_despliegue.html'; out.write_text(construir_html(base, filas, conteo, n), encoding='utf-8'); print(f'HTML: {out}')
    pdf = out.with_suffix('.pdf'); m.a_pdf(out, pdf)
    try:
        import fitz
        doc = fitz.open(pdf); pix = doc[0].get_pixmap(dpi=192); png = base / 'LAMINA_3_despliegue.png'; pix.save(png); print(f'PNG: {png} · {pix.width}×{pix.height}')
        if doc.page_count != 1: print(f'AVISO: el PDF tiene {doc.page_count} páginas (se esperaba 1): la lámina se desbordó.')
    except ImportError:
        print('AVISO: sin PyMuPDF no se genera el PNG (pip install pymupdf)')
    pptx = base / 'LAMINA_3_despliegue.pptx'
    if construir_pptx(base, filas, conteo, n, pptx): print(f'PPTX: {pptx}')

if __name__ == '__main__':
    main()
